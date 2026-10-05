"""Builder generico de planilhas .xlsx para os relatorios (Plan Mode
docs/arquitetura/planmode-relatorios-2026-09-25.md, Bloco 1).

Modulo puro: recebe dado ja formatado, devolve bytes. Sem FastAPI/DB --
quem monta o dado (Bloco 2, services/relatorios.py) decide o que entra em
cada aba; este arquivo so sabe desenhar tabela.
"""

from __future__ import annotations

from collections.abc import Sequence
from datetime import date, datetime
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

_COR_CABECALHO = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
_FONTE_CABECALHO = Font(color="FFFFFF", bold=True)
_FONTE_VARIACAO_POSITIVA = Font(color="15803D")
_FONTE_VARIACAO_NEGATIVA = Font(color="B91C1C")
_LARGURA_MAXIMA_COLUNA = 60

FORMATO_MOEDA = '"R$" #,##0.00'
FORMATO_DATA = "DD/MM/YYYY"


def nova_pasta() -> Workbook:
    pasta = Workbook()
    # Workbook() ja vem com uma aba "Sheet" vazia -- removida pra quem
    # chamar escrever_aba_tabela nao acabar com uma aba fantasma no arquivo.
    pasta.remove(pasta.active)
    return pasta


def escrever_aba_tabela(
    pasta: Workbook,
    titulo: str,
    colunas: Sequence[str],
    linhas: Sequence[Sequence[Any]],
    *,
    colunas_moeda: Sequence[str] = (),
    colunas_data: Sequence[str] = (),
) -> Worksheet:
    """`colunas_moeda`/`colunas_data` (por NOME, não índice -- robusto a
    reordenar coluna no chamador) aplicam `number_format` do Excel, mantendo
    o valor cru (sortável/somável) em vez de string formatada -- a
    formatação pt-BR (`app/reports/formatacao.py`) é só pro Word, que não
    tem número/data nativos numa tabela de texto."""
    aba = pasta.create_sheet(title=titulo[:31])
    aba.append(list(colunas))
    for celula in aba[1]:
        celula.fill = _COR_CABECALHO
        celula.font = _FONTE_CABECALHO
        celula.alignment = Alignment(horizontal="center")
    for linha in linhas:
        aba.append(list(linha))
        # Valores de fontes externas nunca podem virar fórmulas ao abrir o
        # arquivo. Preservamos o texto original (inclusive o prefixo) e
        # forçamos o tipo OOXML de string; números/datas seguem tipados.
        for celula in aba[aba.max_row]:
            if isinstance(celula.value, str) and _texto_parece_formula(celula.value):
                celula.data_type = "s"
    aba.freeze_panes = "A2"
    if linhas:
        aba.auto_filter.ref = f"A1:{get_column_letter(len(colunas))}{len(linhas) + 1}"
    _ajustar_largura_colunas(aba, colunas, linhas)
    _aplicar_formato_por_coluna(aba, colunas, len(linhas), colunas_moeda, FORMATO_MOEDA)
    _aplicar_formato_por_coluna(aba, colunas, len(linhas), colunas_data, FORMATO_DATA)
    aba.print_options.horizontalCentered = True
    aba.sheet_properties.pageSetUpPr.fitToPage = True
    aba.page_setup.fitToWidth = 1
    aba.page_setup.fitToHeight = 0
    aba.print_title_rows = "1:1"
    aba.print_area = f"A1:{get_column_letter(len(colunas))}{max(1, len(linhas) + 1)}" if colunas else None
    return aba


def _texto_parece_formula(valor: str) -> bool:
    """Detecta prefixos interpretáveis como fórmula após espaços/controles."""
    return valor.lstrip().startswith(("=", "+", "-", "@"))


def colorir_variacao(aba: Worksheet, coluna: int, linha: int, valor: float | None) -> None:
    """Pinta a celula de verde/vermelho conforme o sinal de `valor` (variacao
    percentual ano a ano). `valor is None` (primeiro ano da serie) nao pinta."""
    if valor is None:
        return
    celula = aba.cell(row=linha, column=coluna)
    celula.font = _FONTE_VARIACAO_POSITIVA if valor >= 0 else _FONTE_VARIACAO_NEGATIVA


def gerar_bytes(pasta: Workbook) -> bytes:
    buffer = BytesIO()
    pasta.save(buffer)
    return buffer.getvalue()


def _aplicar_formato_por_coluna(
    aba: Worksheet,
    colunas: Sequence[str],
    total_linhas: int,
    nomes_alvo: Sequence[str],
    formato: str,
) -> None:
    if not nomes_alvo or not total_linhas:
        return
    alvo = set(nomes_alvo)
    for indice, cabecalho in enumerate(colunas, start=1):
        if cabecalho not in alvo:
            continue
        for linha_num in range(2, total_linhas + 2):
            aba.cell(row=linha_num, column=indice).number_format = formato


def _ajustar_largura_colunas(aba: Worksheet, colunas: Sequence[str], linhas: Sequence[Sequence[Any]]) -> None:
    for indice, cabecalho in enumerate(colunas, start=1):
        maior = len(str(cabecalho))
        for linha in linhas:
            if indice - 1 < len(linha):
                valor = linha[indice - 1]
                if isinstance(valor, (date, datetime)):
                    tamanho = 10
                elif isinstance(valor, (int, float)):
                    tamanho = len(f"{valor:,.2f}") + 3
                else:
                    tamanho = len(str(valor)) if valor is not None else 0
                maior = max(maior, tamanho)
        aba.column_dimensions[get_column_letter(indice)].width = min(maior + 2, _LARGURA_MAXIMA_COLUNA)
        for celulas in aba.iter_cols(min_col=indice, max_col=indice):
            for celula in celulas:
                celula.alignment = Alignment(vertical="center", wrap_text=True)
