"""DOCX tabular sem reter a árvore XML de todas as linhas em memória."""

from collections.abc import Sequence
from copy import deepcopy
from dataclasses import dataclass
from io import BytesIO
from typing import Any
from zipfile import ZIP_DEFLATED, ZipFile

from docx.document import Document as DocumentType
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.table import Table, _Row
from lxml import etree

from app.reports import docx_builder

_LINHAS_TABELA_CURTA = 6


@dataclass(frozen=True)
class SecaoTabela:
    titulo: str
    colunas: Sequence[str]
    linhas: Sequence[Sequence[Any]]
    pesos_colunas: Sequence[float] | None = None
    fonte_pt: float | None = None


def _escrever_tabela(xml: etree.xmlfile, tabela: Table, secao: SecaoTabela) -> None:
    # Seção vazia ganha uma linha de aviso, igual a `docx_builder.adicionar_tabela`.
    total = max(len(secao.linhas), 1) + 1
    cabecalho = tabela.rows[0]
    docx_builder._formatar_linha_tabela(cabecalho, secao.fonte_pt, indice=0, total=total)
    with xml.element(tabela._tbl.tag, tabela._tbl.attrib):
        xml.write(tabela._tbl.tblPr)
        xml.write(tabela._tbl.tblGrid)
        xml.write(cabecalho._tr)
        modelo = tabela.add_row()
        for celula in modelo.cells:
            # Mesma apresentação de `docx_builder.adicionar_tabela`.
            celula.paragraphs[0].add_run().font.size = Pt(8.5)
            celula.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        docx_builder._formatar_linha_tabela(modelo, secao.fonte_pt, indice=1, total=total)
        tabela._tbl.remove(modelo._tr)
        for indice, valores in enumerate(secao.linhas, start=1):
            # Largura, fonte e margens são constantes: copiar a linha já
            # formatada evita reconstruir essas propriedades para cada célula.
            linha = _Row(deepcopy(modelo._tr), tabela)
            tabela._tbl.append(linha._tr)
            try:
                for celula, valor in zip(linha.cells, valores, strict=True):
                    celula.paragraphs[0].runs[0].text = "" if valor is None else str(valor)
                if total <= 6 and indice == total - 1:
                    docx_builder._formatar_linha_tabela(linha, secao.fonte_pt, indice=indice, total=total)
                xml.write(linha._tr)
            finally:
                # A linha já está no ZIP; não acumular XML ou wrappers por município.
                tabela._tbl.remove(linha._tr)
        if not secao.linhas:
            vazia = _Row(deepcopy(modelo._tr), tabela)
            vazia.cells[0].paragraphs[0].runs[0].text = docx_builder.TEXTO_SEM_REGISTROS
            xml.write(vazia._tr)


def gerar_bytes_tabelados(documento: DocumentType, secoes: Sequence[SecaoTabela]) -> bytes:
    """Preserva o pacote/template; só document.xml é gravado progressivamente.

    Sem arquivos temporários, truncamento de dados ou coleta de lixo forçada.
    A trava cobre também iteração/serialização, não apenas criação de gerador.
    """
    with docx_builder._TRAVA_XML:
        if documento.tables:
            raise ValueError("O documento tabulado deve iniciar sem tabelas.")
        tabelas = []
        for secao in secoes:
            # Título preso (keepNext) a uma tabela longa vai sozinho para a página seguinte.
            docx_builder.adicionar_titulo(documento, secao.titulo, solto=len(secao.linhas) > _LINHAS_TABELA_CURTA)
            tabelas.append(
                docx_builder.adicionar_tabela(
                    documento,
                    secao.colunas,
                    [],
                    pesos_colunas=secao.pesos_colunas,
                    fonte_pt=secao.fonte_pt,
                    texto_vazio=None,
                )
            )
        modelo = docx_builder.gerar_bytes(documento)
        saida = BytesIO()
        with ZipFile(BytesIO(modelo)) as origem, ZipFile(saida, "w", compression=ZIP_DEFLATED) as destino:
            for parte in origem.infolist():
                if parte.filename != "word/document.xml":
                    destino.writestr(parte, origem.read(parte.filename))
                    continue
                with destino.open(parte, "w") as membro, etree.xmlfile(membro, encoding="UTF-8") as xml:
                    xml.write_declaration(standalone=True)
                    raiz = documento._element
                    with xml.element(raiz.tag, raiz.attrib, nsmap=raiz.nsmap):
                        with xml.element(raiz.body.tag, raiz.body.attrib):
                            tabelas_secoes = iter(zip(tabelas, secoes, strict=True))
                            for bloco in raiz.body:
                                if bloco.tag == qn("w:tbl"):
                                    tabela, secao = next(tabelas_secoes)
                                    _escrever_tabela(xml, tabela, secao)
                                else:
                                    xml.write(bloco)
        return saida.getvalue()
