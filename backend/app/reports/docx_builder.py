"""Builder generico de documentos .docx para os relatorios (Plan Mode
docs/arquitetura/planmode-relatorios-2026-09-25.md, Bloco 1).

Modulo puro: recebe dado ja formatado, devolve bytes. Sem FastAPI/DB.
Porta o padrao de estilo ja provado em producao pelo projeto irmao
nota-informativa-decan (RGBColor de variacao, tabela com borda via oxml),
sem importar codigo de la (ver "Achados" do plan-mode -- projetos com
infra/dominio/postura de seguranca diferentes, so a tecnica foi reaproveitada).
"""

from __future__ import annotations

from collections.abc import Callable, Sequence
from datetime import datetime
from functools import wraps
from io import BytesIO
from math import isfinite
from threading import RLock
from typing import Any
from zoneinfo import ZoneInfo

from docx import Document
from docx.document import Document as DocumentType
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.oxml.table import CT_Tbl
from docx.oxml.xmlchemy import BaseOxmlElement
from docx.section import Section
from docx.shared import Cm, Emu, Inches, Pt, RGBColor
from docx.table import Table, _Row
from docx.text.paragraph import Paragraph
from docx.text.run import Run

from app.reports.formatacao import formatar_data, formatar_moeda

_VERDE_VARIACAO = RGBColor(0x15, 0x80, 0x3D)
_VERMELHO_VARIACAO = RGBColor(0xB9, 0x1C, 0x1C)
_CINZA_LEGENDA = RGBColor(0x6B, 0x72, 0x80)
_PRETO = RGBColor(0, 0, 0)

# python-docx compartilha parsers XML entre documentos. A rota síncrona roda
# em threads: proteger as operações evita corrupção de namespace do template.
# Reentrante porque novo_documento chama adicionar_titulo, por exemplo.
_TRAVA_XML = RLock()


def _serializar_xml[**P, R](funcao: Callable[P, R]) -> Callable[P, R]:
    @wraps(funcao)
    def executar(*args: P.args, **kwargs: P.kwargs) -> R:
        with _TRAVA_XML:
            return funcao(*args, **kwargs)

    return executar


# Mesmo timbre institucional usado nos briefings reais do departamento
# (ver docs/relatorios em `data/relatorios/`, anexados pelo usuário
# 2026-09-26) -- documento gerado por aqui deixa explícito que é
# automático, nunca se apresenta como Nota Informativa/Briefing já
# revisado por um analista.
_CABECALHO_INSTITUCIONAL = (
    "Ministério da Saúde",
    "SAES -- Secretaria de Atenção Especializada à Saúde",
    "DECAN -- Departamento de Atenção ao Câncer",
    "CGPCAN -- Coordenação-Geral de Planejamento e Monitoramento do Câncer",
)


def _inserir_bloco(documento: DocumentType, bloco: BaseOxmlElement) -> None:
    # O corpo termina em sectPr. A biblioteca procura esse filho desde o
    # início a cada inserção, tornando blocos narrativos longos quadráticos.
    # Inserção no mesmo lugar, sem remover a seção ou cachear sua geometria.
    corpo = documento._element.body
    ultimo = next(corpo.iterchildren(reversed=True), None)
    if ultimo is not None and ultimo.tag == qn("w:sectPr"):
        ultimo.addprevious(bloco)
    else:
        corpo.append(bloco)


def _novo_paragrafo(documento: DocumentType, texto: str = "", *, estilo: str | None = None) -> Paragraph:
    elemento = OxmlElement("w:p")
    _inserir_bloco(documento, elemento)
    paragrafo = Paragraph(elemento, documento._body)
    if texto:
        paragrafo.add_run(texto)
    if estilo is not None:
        paragrafo.style = estilo
    return paragrafo


def _ultima_secao(documento: DocumentType) -> Section:
    corpo = documento._element.body
    ultimo = next(corpo.iterchildren(reversed=True), None)
    if ultimo is not None and ultimo.tag == qn("w:sectPr"):
        return Section(ultimo, documento.part)
    return documento.sections[-1]


def _nova_tabela(documento: DocumentType, colunas: int) -> Table:
    secao = _ultima_secao(documento)
    largura = (secao.page_width or Inches(8.5)) - (secao.left_margin or Inches(1)) - (secao.right_margin or Inches(1))
    elemento = CT_Tbl.new_tbl(1, colunas, Emu(largura))
    _inserir_bloco(documento, elemento)
    tabela = Table(elemento, documento._body)
    tabela.style = None
    return tabela


@_serializar_xml
def novo_documento(titulo: str, *, paisagem: bool = False) -> DocumentType:
    documento = Document()
    secao = documento.sections[0]
    if paisagem:
        # Tabelas largas (muitas colunas) cabem melhor na página deitada.
        secao.orientation = WD_ORIENT.LANDSCAPE
        secao.page_width = Cm(29.7)
        secao.page_height = Cm(21)
    else:
        secao.page_width = Cm(21)
        secao.page_height = Cm(29.7)
    secao.top_margin = secao.bottom_margin = Cm(2)
    secao.left_margin = secao.right_margin = Cm(2)
    estilos = documento.styles
    normal = estilos["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = _PRETO
    normal.paragraph_format.space_after = Pt(4)
    for nome, tamanho in (("Title", 16), ("Heading 1", 12), ("Heading 2", 10.5), ("Heading 3", 10), ("Heading 4", 9.5)):
        estilo = estilos[nome]
        estilo.font.name = "Arial"
        estilo.font.size = Pt(tamanho)
        estilo.font.color.rgb = _PRETO
        estilo.font.bold = True
        estilo.paragraph_format.space_before = Pt(9 if nome != "Title" else 0)
        estilo.paragraph_format.space_after = Pt(4)
        estilo.paragraph_format.keep_with_next = True
        estilo.paragraph_format.keep_together = True
        # O estilo Title do Word pode carregar uma borda azul temática.
        ppr = estilo.element.pPr
        if ppr is not None:
            bordas = ppr.find(qn("w:pBdr"))
            if bordas is not None:
                ppr.remove(bordas)
    rodape = secao.footer.paragraphs[0]
    rodape.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    texto = rodape.add_run("SIGEO  ·  Página ")
    texto.font.size = Pt(8)
    texto.font.color.rgb = _CINZA_LEGENDA
    campo = OxmlElement("w:fldSimple")
    campo.set(qn("w:instr"), "PAGE")
    rodape._p.append(campo)
    adicionar_titulo(documento, titulo, nivel=0)
    return documento


@_serializar_xml
def adicionar_cabecalho_institucional(documento: DocumentType, *, gerado_em: datetime) -> None:
    for linha in _CABECALHO_INSTITUCIONAL:
        paragrafo = _novo_paragrafo(documento)
        run = paragrafo.add_run(linha)
        run.bold = linha == _CABECALHO_INSTITUCIONAL[0]
        run.font.size = Pt(10)
    aviso = _novo_paragrafo(documento)
    run_aviso = aviso.add_run(
        f"Documento gerado automaticamente pelo SIGEO em {gerado_em.astimezone(ZoneInfo('America/Sao_Paulo')).strftime('%d/%m/%Y %H:%M')} "
        "(horário de Brasília) "
        "-- não substitui conferência humana antes de circulação oficial."
    )
    run_aviso.italic = True
    run_aviso.font.size = Pt(9)
    run_aviso.font.color.rgb = _CINZA_LEGENDA


@_serializar_xml
def adicionar_titulo(documento: DocumentType, texto: str, *, nivel: int = 1) -> None:
    if not 0 <= nivel <= 9:
        raise ValueError(f"level must be in range 0-9, got {nivel}")
    _novo_paragrafo(documento, texto, estilo="Title" if nivel == 0 else f"Heading {nivel}")


@_serializar_xml
def adicionar_paragrafo(documento: DocumentType, texto: str) -> None:
    _novo_paragrafo(documento, texto)


@_serializar_xml
def adicionar_paragrafo_rotulo_valor(documento: DocumentType, rotulo: str, valor: str) -> None:
    """`**Rótulo:** valor` -- mesmo padrão usado nos briefings reais do
    departamento pra ficha de instrumento (Convênio/SEI/Objeto/...)."""
    paragrafo = _novo_paragrafo(documento)
    run_rotulo = paragrafo.add_run(f"{rotulo}: ")
    run_rotulo.bold = True
    paragrafo.add_run(valor)


@_serializar_xml
def adicionar_legenda_tabela(documento: DocumentType, numero: int, texto: str) -> None:
    paragrafo = _novo_paragrafo(documento)
    run = paragrafo.add_run(f"Tabela {numero}. {texto}")
    run.italic = True


@_serializar_xml
def adicionar_fonte(documento: DocumentType, texto: str) -> None:
    paragrafo = _novo_paragrafo(documento)
    run = paragrafo.add_run(f"Fonte: {texto}")
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = _CINZA_LEGENDA


TEXTO_SEM_REGISTROS = "Nenhum registro neste recorte."


@_serializar_xml
def adicionar_tabela(
    documento: DocumentType,
    colunas: Sequence[str],
    linhas: Sequence[Sequence[Any]],
    *,
    pesos_colunas: Sequence[float] | None = None,
    fonte_pt: float | None = None,
    texto_vazio: str | None = TEXTO_SEM_REGISTROS,
) -> Table:
    if pesos_colunas is not None and (
        len(pesos_colunas) != len(colunas) or any(not isfinite(p) or p <= 0 for p in pesos_colunas)
    ):
        raise ValueError("Informe um peso positivo e finito para cada coluna.")
    tabela = _nova_tabela(documento, len(colunas))
    tabela.alignment = WD_TABLE_ALIGNMENT.CENTER
    tabela.style = "Table Grid"
    _aplicar_bordas(tabela)
    propriedade_header = OxmlElement("w:tblHeader")
    propriedade_header.set(qn("w:val"), "true")
    tabela.rows[0]._tr.get_or_add_trPr().append(propriedade_header)
    for celula, cabecalho in zip(tabela.rows[0].cells, colunas, strict=True):
        run = celula.paragraphs[0].add_run(cabecalho)
        run.bold = True
        run.font.size = Pt(8.5)
        celula.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for linha in linhas:
        row = tabela.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        celulas = row.cells
        for celula, valor in zip(celulas, linha, strict=True):
            run = celula.paragraphs[0].add_run("" if valor is None else str(valor))
            run.font.size = Pt(8.5)
            celula.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if not linhas and texto_vazio:
        celulas = tabela.add_row().cells
        for posicao, celula in enumerate(celulas):
            celula.paragraphs[0].add_run(texto_vazio if posicao == 0 else "").font.size = Pt(8.5)
            celula.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if pesos_colunas is not None:
        _aplicar_larguras(documento, tabela, pesos_colunas)
    # A tabela já está completa; len(rows) refaz a lista de elementos XML.
    total_linhas = len(tabela.rows)
    for indice, linha_tabela in enumerate(tabela.rows):
        _formatar_linha_tabela(linha_tabela, fonte_pt, indice=indice, total=total_linhas)
    return tabela


def _formatar_linha_tabela(linha: _Row, fonte_pt: float | None, *, indice: int, total: int) -> None:
    propriedades = linha._tr.get_or_add_trPr()
    if propriedades.find(qn("w:cantSplit")) is None:
        propriedades.append(OxmlElement("w:cantSplit"))
    # A mesma paginação vale para a árvore completa e para a gravação progressiva.
    manter_junta = indice < total - 1 and (total <= 6 or indice == 0)
    for celula in linha.cells:
        for paragrafo in celula.paragraphs:
            paragrafo.paragraph_format.keep_with_next = True if manter_junta else None
            if fonte_pt is not None:
                paragrafo.paragraph_format.space_before = Pt(0)
                paragrafo.paragraph_format.space_after = Pt(2)
                for run in paragrafo.runs:
                    run.font.size = Pt(fonte_pt)


def _aplicar_larguras(documento: DocumentType, tabela: Table, pesos: Sequence[float]) -> None:
    secao = _ultima_secao(documento)
    largura = int(secao.page_width or 0) - int(secao.left_margin or 0) - int(secao.right_margin or 0)
    larguras = [Emu(round(largura * peso / sum(pesos))) for peso in pesos]
    tabela.autofit = False
    for coluna, medida in zip(tabela.columns, larguras, strict=True):
        coluna.width = medida
    for linha in tabela.rows:
        for celula, medida in zip(linha.cells, larguras, strict=True):
            celula.width = medida
    margens = OxmlElement("w:tblCellMar")
    for lado in ("left", "right"):
        margem = OxmlElement(f"w:{lado}")
        margem.set(qn("w:w"), "45")
        margem.set(qn("w:type"), "dxa")
        margens.append(margem)
    tabela._tbl.tblPr.append(margens)


def linhas_como_texto(
    colunas: Sequence[str],
    linhas: Sequence[Sequence[Any]],
    *,
    colunas_moeda: Sequence[str] = (),
    colunas_data: Sequence[str] = (),
    colunas_inteiro: Sequence[str] = (),
) -> list[list[str]]:
    """Word não tem número/data nativos numa tabela de texto -- formata pt-BR
    (`app/reports/formatacao.py`) por NOME de coluna antes de passar pra
    `adicionar_tabela`. Complementa `xlsx_builder.escrever_aba_tabela`, que
    faz o equivalente no Excel via `number_format` (mantendo o valor cru)."""
    moeda = set(colunas_moeda)
    data = set(colunas_data)
    inteiro = set(colunas_inteiro)
    resultado: list[list[str]] = []
    for linha in linhas:
        nova_linha: list[str] = []
        for cabecalho, valor in zip(colunas, linha, strict=True):
            if cabecalho in moeda:
                nova_linha.append(formatar_moeda(valor))
            elif cabecalho in data:
                nova_linha.append(formatar_data(valor))
            elif cabecalho in inteiro and valor is not None:
                nova_linha.append(f"{int(valor):,}".replace(",", "."))
            else:
                nova_linha.append("—" if valor is None else str(valor))
        resultado.append(nova_linha)
    return resultado


@_serializar_xml
def escrever_variacao(paragrafo: Paragraph, valor: float | None) -> Run:
    """Adiciona ao paragrafo o texto da variacao percentual, colorido de
    verde (positiva) ou vermelho (negativa). `valor is None` (primeiro ano
    da serie) escreve "--" sem cor."""
    if valor is None:
        return paragrafo.add_run("--")
    run = paragrafo.add_run(f"{valor:+.1f}%")
    run.font.color.rgb = _VERDE_VARIACAO if valor >= 0 else _VERMELHO_VARIACAO
    return run


@_serializar_xml
def gerar_bytes(documento: DocumentType) -> bytes:
    buffer = BytesIO()
    documento.save(buffer)
    return buffer.getvalue()


def _aplicar_bordas(tabela: Table) -> None:
    propriedades = tabela._tbl.tblPr
    bordas = OxmlElement("w:tblBorders")
    for nome in ("top", "left", "bottom", "right", "insideH", "insideV"):
        borda = OxmlElement(f"w:{nome}")
        borda.set(qn("w:val"), "single")
        borda.set(qn("w:sz"), "4")
        borda.set(qn("w:color"), "999999")
        bordas.append(borda)
    propriedades.append(bordas)
