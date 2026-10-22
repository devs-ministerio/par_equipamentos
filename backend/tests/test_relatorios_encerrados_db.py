"""SQL real, HTTP autenticado e arquivos produzidos pelo mesmo classificador."""

from datetime import date, datetime, timezone
from io import BytesIO

import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select

from app.db.base import SessionLocal, get_db
from app.db.models import Convenio, EventoMarco, InstrumentoEquipamento, MarcoCatalogo, PropostaCandidata
from app.main import app
from app.schemas_relatorios_encerrados import FiltroEncerramento
from app.services.relatorios import FiltroRelatorio, montar_relatorio
from app.services.relatorios_encerrados import extrair_encerrados

PREFIXO = "Encerrados sintéticos 99081"
FECHADA = "Prestação de Contas Concluída"


@pytest.fixture
def carteira_encerrada(monkeypatch):
    with SessionLocal() as db:
        marco = db.scalar(select(MarcoCatalogo).where(MarcoCatalogo.codigo == "fase_concluido"))
        assert marco is not None
        for indice in range(1, 8):
            numero = f"990810{indice}"
            db.add(
                Convenio(
                    numero=numero,
                    convenente_nome=f"{PREFIXO} {indice}",
                    tipo_contratacao="Convênio",
                    uf="BA",
                    municipio="Anápolis",
                    ano_instrumento=2025,
                    valor_global=0 if indice == 1 else 100,
                    situacao=FECHADA if indice in (1, 2, 5) else "Em execução",
                    fora_escopo_equipamentos=indice == 7,
                )
            )
            if indice == 1:
                continue
            inst = InstrumentoEquipamento(
                nr_convenio=numero,
                nome_convenente=f"{PREFIXO} {indice}",
                tipo_contratacao="Convênio",
                uf="BA",
                municipio="Anápolis",
                ano_instrumento=2025,
                situacao_prestacao_contas="Em execução" if indice == 5 else None,
            )
            db.add(inst)
            db.flush()
            db.add(
                EventoMarco(
                    instrumento_id=inst.id,
                    marco_id=marco.id,
                    data_ocorrencia=None if indice in (2, 4) else date(2025, 1, 1),
                    deletado_em=datetime.now(timezone.utc) if indice == 6 else None,
                    motivo_exclusao="Fixture histórica excluída" if indice == 6 else None,
                )
            )
        db.add(
            PropostaCandidata(
                id_proposta=9908190,
                cd_parceria="20259908190",
                tem_parceria=True,
                cnpj_ente_recebedor="00000000000001",
                nm_proponente=f"{PREFIXO} parceria",
                ds_objeto="Equipamento sintético",
                nm_programa="Programa sintético",
                componente_batido="Componente sintético",
                id_programa=1,
                uf="BA",
                data_proposta=date(2025, 1, 1),
                situacao_proposta="Aprovada",
                metas_resumo={
                    "parceria": {"cd_parceria": "20259908190"},
                    "timeline_financeira": {
                        "ordens_pagamento": [{"id_ordem_pagamento": 9908191, "in_situacao_op": "Paga"}]
                    },
                },
            )
        )
        db.flush()
        monkeypatch.setitem(app.dependency_overrides, get_db, lambda: db)
        yield db
        db.rollback()


def test_reconcilia_escopo_eventos_ativos_e_novos_convenios_sem_monitoramento(carteira_encerrada):
    # Given / When
    resultado = extrair_encerrados(db=carteira_encerrada, filtro=FiltroEncerramento(busca=PREFIXO))
    # Then
    assert [i.numero for i in resultado.itens] == ["9908101", "9908102", "9908103", "20259908190"]
    assert {i.numero for i in resultado.pendencias} == {"9908104", "9908105"}
    assert resultado.itens[0].valor_global == 0
    assert resultado.itens[2].motivo == "Concluído internamente — prestação de contas pendente"


def test_instrumento_sem_tipo_contratacao_nao_derruba_a_extracao(carteira_encerrada):
    # Given: coluna nullable; instrumento legado sem tipo e fora de `convenio`
    carteira_encerrada.add(
        InstrumentoEquipamento(nr_convenio="9908199", nome_convenente=f"{PREFIXO} sem tipo", tipo_contratacao=None)
    )
    carteira_encerrada.flush()
    # When
    resultado = extrair_encerrados(db=carteira_encerrada, filtro=FiltroEncerramento(busca=PREFIXO))
    # Then: sem erro de validacao; tratado como "Convênio", igual ao ramo de `convenio`
    assert "9908199" not in {i.numero for i in resultado.itens}


@pytest.mark.parametrize("nivel", ["simplificado", "completo"])
def test_previa_http_e_arquivos_tem_mesmas_identidades_e_motivos(carteira_encerrada, headers_autenticados, nivel):
    # Given
    client = TestClient(app)
    filtro = FiltroRelatorio(escopo="uf", uf="BA", busca=PREFIXO, encerramento="encerrados")
    # When
    previa = client.get(
        "/relatorios/encerrados", params={"escopo": "uf", "uf": "BA", "busca": PREFIXO}, headers=headers_autenticados
    )
    assert previa.status_code == 200, previa.text
    resposta = previa.json()
    planilha = load_workbook(
        BytesIO(
            montar_relatorio(
                db=carteira_encerrada, formato="xlsx", nivel=nivel, tipo_relatorio="instrumentos_repasse", filtro=filtro
            )
        )
    )
    documento = Document(
        BytesIO(
            montar_relatorio(
                db=carteira_encerrada, formato="docx", nivel=nivel, tipo_relatorio="instrumentos_repasse", filtro=filtro
            )
        )
    )
    # Then
    numeros = [i["numero"] for i in resposta["itens"]]
    assert numeros == [r[0] for r in planilha["Encerrados"].iter_rows(min_row=2, values_only=True)]
    assert numeros == [r.cells[0].text for r in documento.tables[0].rows[1:]]
    assert [i["motivo"] for i in resposta["itens"]] == [
        r[7] for r in planilha["Encerrados"].iter_rows(min_row=2, values_only=True)
    ]
    assert resposta["total_pendencias"] == planilha["Pendências"].max_row - 1 == 2
    assert [r.cells[0].text for r in documento.tables[1].rows[1:]] == [i["numero"] for i in resposta["pendencias"]]


def test_contrato_exige_sessao_e_valida_paginacao(carteira_encerrada, headers_autenticados):
    # Given
    client = TestClient(app)
    # When / Then
    assert client.get("/relatorios/encerrados").status_code == 401
    for params in ({"pagina": 0}, {"limite": 101}, {"escopo": "uf"}, {"encerramento": "outro"}):
        assert client.get("/relatorios/encerrados", params=params, headers=headers_autenticados).status_code == 422
    resposta = client.get(
        "/relatorios/encerrados",
        params={"escopo": "brasil", "busca": PREFIXO, "limite": 1, "pagina": 2},
        headers=headers_autenticados,
    )
    assert resposta.status_code == 200, resposta.text
    pagina = resposta.json()
    assert pagina["total"] == 4 and pagina["itens"][0]["numero"] == "9908102"
    assert (
        client.get(
            "/relatorios", params={"formato": "xlsx", "encerramento": "outro"}, headers=headers_autenticados
        ).status_code
        == 422
    )
