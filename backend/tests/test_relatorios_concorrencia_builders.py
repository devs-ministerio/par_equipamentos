"""Sobreposição de gerações reais, sem SQL, rede ou artefato de usuário.

Não é um benchmark de capacidade HTTP/SQL nem aceite visual de Word/Excel.
Cada worker cria sua própria pasta/documento e valida os bytes que recebeu.
"""

from concurrent.futures import ThreadPoolExecutor
from datetime import date, datetime
from io import BytesIO
from threading import Barrier
from zipfile import ZipFile

import pytest
from docx import Document
from docx.oxml.ns import qn
from openpyxl import load_workbook

from app.reports import docx_builder, xlsx_builder

_COLUNAS = ["Identificador", "Nome", "Data", "Valor"]


def _linhas_sinteticas(lote: int, total: int) -> list[list[str | date | int | None]]:
    return [
        [
            f"000{lote:02d}{indice:06d}",
            f"Solicitação {lote} · Unidade sintética {indice}",
            date(2026, 10, 20),
            None if indice % 5 == 4 else indice,
        ]
        for indice in range(total)
    ]


def _gerar_xlsx(lote: int, total: int) -> bytes:
    linhas = _linhas_sinteticas(lote, total)
    pasta = xlsx_builder.nova_pasta()
    xlsx_builder.escrever_aba_tabela(
        pasta,
        f"Solicitação {lote}",
        _COLUNAS,
        linhas,
        colunas_moeda=["Valor"],
        colunas_data=["Data"],
    )
    return xlsx_builder.gerar_bytes(pasta)


def _gerar_docx(lote: int, total: int) -> bytes:
    documento = docx_builder.novo_documento(f"Solicitação {lote}")
    docx_builder.adicionar_tabela(documento, _COLUNAS, _linhas_sinteticas(lote, total))
    return docx_builder.gerar_bytes(documento)


def _conferir_xlsx(conteudo: bytes, lote: int, total: int) -> None:
    pasta = load_workbook(BytesIO(conteudo))
    try:
        assert pasta.sheetnames == [f"Solicitação {lote}"]
        aba = pasta.active
        assert aba is not None
        assert aba.max_row == total + 1
        assert aba.freeze_panes == "A2"
        assert aba.auto_filter.ref == f"A1:D{total + 1}"
        assert [celula.value for celula in aba[1]] == _COLUNAS
        for indice, linha in enumerate(aba.iter_rows(min_row=2)):
            assert linha[0].value == f"000{lote:02d}{indice:06d}"
            assert linha[0].data_type == "s"
            assert linha[1].value == f"Solicitação {lote} · Unidade sintética {indice}"
            assert linha[2].value == datetime(2026, 10, 20)
            assert linha[2].number_format == xlsx_builder.FORMATO_DATA
            assert linha[3].value == (None if indice % 5 == 4 else indice)
            assert linha[3].number_format == xlsx_builder.FORMATO_MOEDA
            if indice == 0:
                assert linha[3].value == 0
    finally:
        pasta.close()


def _conferir_docx(conteudo: bytes, lote: int, total: int) -> None:
    documento = Document(BytesIO(conteudo))
    assert documento.paragraphs[0].text == f"Solicitação {lote}"
    assert len(documento.tables) == 1
    tabela = documento.tables[0]
    assert len(tabela.rows) == total + 1
    assert [celula.text for celula in tabela.rows[0].cells] == _COLUNAS
    assert tabela.rows[0]._tr.find(f"{qn('w:trPr')}/{qn('w:tblHeader')}") is not None
    # Percorrer as células uma vez evita repetir a expansão da tabela por linha.
    celulas = tabela._cells
    for indice, valores in enumerate(_linhas_sinteticas(lote, total), start=1):
        assert [celula.text for celula in celulas[indice * 4 : (indice + 1) * 4]] == [
            "" if valor is None else str(valor) for valor in valores
        ]
        assert tabela.rows[indice]._tr.find(f"{qn('w:trPr')}/{qn('w:cantSplit')}") is not None


def _solicitacao(formato: str, lote: int, total: int, inicio: Barrier) -> bytes:
    inicio.wait(timeout=30)
    return _gerar_xlsx(lote, total) if formato == "xlsx" else _gerar_docx(lote, total)


@pytest.mark.parametrize(
    "pedidos",
    [
        [("xlsx", 5000)] * 4,
        [("docx", 500)] * 4,
        [("xlsx", 5000), ("docx", 500), ("xlsx", 21), ("docx", 21)],
    ],
    ids=["quatro-xlsx-5000", "quatro-docx-500", "formatos-e-volumes-mistos"],
)
def test_geracoes_sobrepostas_preservam_conteudo_e_formato_de_cada_solicitacao(pedidos):
    # Given
    inicio = Barrier(len(pedidos))
    with ThreadPoolExecutor(max_workers=len(pedidos)) as executor:
        # When
        tarefas = [
            executor.submit(_solicitacao, formato, lote, total, inicio)
            for lote, (formato, total) in enumerate(pedidos, start=1)
        ]
        resultados = [tarefa.result(timeout=180) for tarefa in tarefas]
    # Then
    assert len(set(resultados)) == len(pedidos)
    for lote, ((formato, total), conteudo) in enumerate(zip(pedidos, resultados, strict=True), start=1):
        with ZipFile(BytesIO(conteudo)) as pacote:
            assert pacote.testzip() is None
        if formato == "xlsx":
            _conferir_xlsx(conteudo, lote, total)
        else:
            _conferir_docx(conteudo, lote, total)


def test_erro_de_entrada_nao_bloqueia_geracao_docx_de_outra_thread():
    # Given
    documento = docx_builder.novo_documento("Entrada inválida")
    # When
    with pytest.raises(ValueError, match="peso positivo"):
        docx_builder.adicionar_tabela(documento, ["Coluna"], [], pesos_colunas=[0])
    with ThreadPoolExecutor(max_workers=1) as executor:
        conteudo = executor.submit(_gerar_docx, 9, 1).result(timeout=30)
    # Then
    assert len(documento.tables) == 0
    _conferir_docx(conteudo, 9, 1)
