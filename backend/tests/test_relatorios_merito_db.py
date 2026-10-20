"""Conteúdo dos arquivos de mérito com SQL e seleção de publicação reais."""

from io import BytesIO

import pytest
from docx import Document
from openpyxl import load_workbook
from sqlalchemy import update

from app.db.base import SessionLocal
from app.db.models import Competency
from app.domain_errors import ValidationError
from app.services.relatorios import FiltroRelatorio, montar_relatorio
from tests.fixtures_relatorios_merito import FAMILIAS, criar_publicacoes_merito, criar_resultado

CASOS = (
    ("brasil", FiltroRelatorio(escopo="brasil"), (0, 1, 2)),
    ("nordeste", FiltroRelatorio(escopo="regiao", regiao="Nordeste"), (0, 1)),
    ("bahia", FiltroRelatorio(escopo="uf", uf="BA"), (0,)),
    ("sao_paulo", FiltroRelatorio(escopo="uf", uf="SP"), (2,)),
    ("vazio", FiltroRelatorio(escopo="uf", uf="AC"), ()),
)


@pytest.fixture
def publicacoes_merito():
    with SessionLocal() as db:
        atuais = criar_publicacoes_merito(db)
        try:
            yield db, atuais
        finally:
            db.rollback()


def ler_secoes(conteudo, formato):
    if formato == "xlsx":
        pasta = load_workbook(BytesIO(conteudo))
        return {aba.title: [list(row) for row in aba.iter_rows(values_only=True)] for aba in pasta}
    documento = Document(BytesIO(conteudo))
    titulos = [p.text for p in documento.paragraphs if p.text.startswith("Cobertura")]
    return {
        titulo: [[c.text for c in row.cells] for row in tabela.rows]
        for titulo, tabela in zip(titulos, documento.tables, strict=True)
    }


def linhas_esperadas(familia, nivel):
    nomes = (
        ["Macro sintética Bahia", "Macro sintética Pernambuco", "Macro sintética São Paulo"]
        if nivel == "simplificado"
        else ["SALVADOR", "RECIFE", "SAO PAULO"]
    )
    oferta, saldo = (1, -1) if familia == "TOMOGRAFO" else (0, -2)
    return [
        [nomes[0], "BA", 100_000, 2, oferta, saldo],
        [nomes[1], "PE", 0, 0, 0, 0],
        [nomes[2], "SP", None, None, None, None],
    ]


def normalizar_word(linhas):
    return [["—" if valor is None else str(valor) for valor in linha] for linha in linhas]


@pytest.mark.parametrize("formato", ["xlsx", "docx"])
@pytest.mark.parametrize("nivel", ["simplificado", "completo"])
@pytest.mark.parametrize("nome,filtro,indices", CASOS, ids=[c[0] for c in CASOS])
def test_merito_exporta_resultados_da_publicacao_correta_por_familia(
    publicacoes_merito, formato, nivel, nome, filtro, indices
):
    # Given: duas famílias com publicação própria, anterior, avulsa e rascunho.
    db, _ = publicacoes_merito
    # When: a cadeia de SQL até os bytes finais roda sem mocks de domínio.
    conteudo = montar_relatorio(db=db, formato=formato, nivel=nivel, tipo_relatorio="analise_merito", filtro=filtro)
    secoes = ler_secoes(conteudo, formato)
    # Then: valores exatos, não só presença das abas nem aprovação por subconjunto.
    for familia in FAMILIAS:
        linhas = secoes[f"Cobertura {familia}"]
        assert linhas[0] == [
            "Macrorregião" if nivel == "simplificado" else "Município",
            "UF",
            "População",
            "Necessário",
            "Disponível SUS",
            "Saldo",
        ]
        esperadas = [linhas_esperadas(familia, nivel)[i] for i in indices]
        assert linhas[1:] == (normalizar_word(esperadas) if formato == "docx" else esperadas)


@pytest.mark.parametrize("formato", ["xlsx", "docx"])
def test_merito_municipio_com_acento_exporta_somente_o_municipio_pedido(publicacoes_merito, formato):
    # Given
    db, _ = publicacoes_merito
    # When
    secoes = ler_secoes(
        montar_relatorio(
            db=db,
            formato=formato,
            nivel="completo",
            tipo_relatorio="analise_merito",
            filtro=FiltroRelatorio(escopo="municipio", uf="SP", municipio="São Paulo"),
        ),
        formato,
    )
    # Then
    for familia in FAMILIAS:
        esperado = [linhas_esperadas(familia, "completo")[2]]
        assert secoes[f"Cobertura {familia}"][1:] == (normalizar_word(esperado) if formato == "docx" else esperado)


@pytest.mark.parametrize("formato", ["xlsx", "docx"])
@pytest.mark.parametrize("nivel", ["simplificado", "completo"])
def test_merito_sem_publicacao_emite_aviso_com_consulta_real(publicacoes_merito, formato, nivel):
    # Given: ponteiros removidos somente na transação sintética revertida ao final.
    db, _ = publicacoes_merito
    db.execute(update(Competency).values(published_execution_id=None))
    # When
    secoes = ler_secoes(
        montar_relatorio(
            db=db,
            formato=formato,
            nivel=nivel,
            tipo_relatorio="analise_merito",
            filtro=FiltroRelatorio(escopo="brasil"),
        ),
        formato,
    )
    # Then
    assert secoes == {
        "Cobertura": [["Aviso"], ["Nenhuma família de equipamento com execução publicada para este filtro."]]
    }


@pytest.mark.parametrize("formato", ["xlsx", "docx"])
@pytest.mark.parametrize("nivel", ["simplificado", "completo"])
def test_merito_recusa_granularidade_cnes_sem_exportar_dado_nacional(publicacoes_merito, formato, nivel):
    # Given
    db, _ = publicacoes_merito
    # When / Then
    with pytest.raises(ValidationError, match="não tem granularidade por estabelecimento"):
        montar_relatorio(
            db=db,
            formato=formato,
            nivel=nivel,
            tipo_relatorio="analise_merito",
            filtro=FiltroRelatorio(escopo="cnes", cnes="0000799"),
        )


@pytest.mark.parametrize("formato", ["xlsx", "docx"])
@pytest.mark.parametrize("nivel", ["simplificado", "completo"])
def test_merito_nao_trunca_na_paginacao_da_interface(publicacoes_merito, formato, nivel, tmp_path):
    # Given: 204 resultados, acima das páginas de 20/50 itens da interface.
    db, atuais = publicacoes_merito
    for indice in range(3, 204):
        criar_resultado(db, atuais["TOMOGRAFO"], "TOMOGRAFO", indice, nome_extra=f" Lote {indice:04}")
    db.flush()
    # When
    conteudo = montar_relatorio(
        db=db,
        formato=formato,
        nivel=nivel,
        tipo_relatorio="analise_merito",
        filtro=FiltroRelatorio(escopo="brasil"),
    )
    secoes = ler_secoes(conteudo, formato)
    linhas = secoes["Cobertura TOMOGRAFO"][1:]
    # Then: nenhuma linha repetida ou omitida, incluindo o fim do lote.
    assert len(linhas) == 204
    assert len({linha[0] for linha in linhas}) == 204
    assert sum(" Lote " in linha[0] for linha in linhas) == 201
    assert any(linha[0].endswith("Lote 0203") for linha in linhas)
    if nivel == "completo":
        (tmp_path / f"merito-longo.{formato}").write_bytes(conteudo)
