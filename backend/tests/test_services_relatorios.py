"""app/services/relatorios.py -- FiltroRelatorio e montar_relatorio, os 2
relatórios (`instrumentos_repasse`/`analise_merito`) (Plan Mode
docs/arquitetura/planmode-relatorios-2026-09-25.md, Blocos 2, 4 e 5)."""

from __future__ import annotations

from datetime import date
from io import BytesIO
from uuid import uuid4

import pytest
from docx import Document
from openpyxl import load_workbook

from app.db.base import SessionLocal
from app.db.models import (
    AcaoMonitoramento,
    Convenio,
    EventoMarco,
    InstrumentoEquipamento,
    MarcoCatalogo,
    PropostaCandidata,
)
from app.domain_errors import ValidationError
from app.reports.formatacao import formatar_data
from app.services.relatorios import (
    FiltroRelatorio,
    _assunto,
    _data_inauguracao_previsao,
    _formatar_ultima_acao,
    _formatar_ultimo_evento,
    _linhas_convenio,
    _linhas_proposta,
    _observacao_publica,
    _resumo_equipamentos,
    _valor_global_confiavel,
    montar_relatorio,
)


def _garantir_convenio_948686(db) -> None:
    """O bloco narrativo de "Informações Detalhadas" itera sobre `Convenio`
    e anexa o `InstrumentoEquipamento` correspondente por `nr_convenio`
    (`_montar_instrumentos_repasse_docx::instrumentos_por_nr`) --
    `scripts/seed_monitoramento.py` só semeia o instrumento monitorado (+
    CNES), nunca o `Convenio` companheiro, então sem isto o bloco do
    948686 nunca apareceria no documento."""
    if db.query(Convenio).filter_by(numero="948686").one_or_none() is not None:
        return
    db.add(
        Convenio(
            numero="948686",
            convenente_nome="INSTITUTO DE GESTAO ESTRATEGICA DE SAUDE DO DISTRITO FEDERAL - IGESDF",
            uf="DF",
            tipo_contratacao="FAF",
        )
    )
    db.flush()


def test_filtro_regiao_sem_valor_e_rejeitado():
    with pytest.raises(ValidationError):
        FiltroRelatorio(escopo="regiao")


def test_filtro_regiao_desconhecida_e_rejeitado():
    with pytest.raises(ValidationError):
        FiltroRelatorio(escopo="regiao", regiao="Atlantida")


def test_filtro_uf_sem_valor_e_rejeitado():
    with pytest.raises(ValidationError):
        FiltroRelatorio(escopo="uf")


def test_filtro_municipio_sem_uf_e_rejeitado():
    with pytest.raises(ValidationError):
        FiltroRelatorio(escopo="municipio", municipio="Cidade Alfa")


def test_filtro_cnes_sem_valor_e_rejeitado():
    with pytest.raises(ValidationError):
        FiltroRelatorio(escopo="cnes")


def test_filtro_brasil_nao_exige_nada():
    filtro = FiltroRelatorio(escopo="brasil")
    assert filtro.ufs() is None
    assert filtro.titulo() == "Brasil"


def test_filtro_regiao_resolve_ufs_da_regiao():
    filtro = FiltroRelatorio(escopo="regiao", regiao="Centro-Oeste")
    ufs = filtro.ufs()
    assert ufs is not None and "DF" in ufs


def test_filtro_cnes_nao_filtra_por_uf():
    filtro = FiltroRelatorio(escopo="cnes", cnes="9000001")
    assert filtro.ufs() is None
    assert filtro.titulo() == "CNES 9000001"


def test_filtro_titulo_inclui_ano_quando_informado():
    filtro = FiltroRelatorio(escopo="uf", uf="SP", ano_inicio=2023)
    assert filtro.titulo() == "UF SP -- 2023"


def test_filtro_titulo_inclui_periodo_quando_ano_inicio_e_fim_diferem():
    filtro = FiltroRelatorio(escopo="uf", uf="SP", ano_inicio=2020, ano_fim=2023)
    assert filtro.titulo() == "UF SP -- 2020 a 2023"


def test_filtro_ano_fim_antes_de_ano_inicio_levanta_erro():
    with pytest.raises(ValidationError):
        FiltroRelatorio(escopo="brasil", ano_inicio=2023, ano_fim=2020)


def test_filtro_ano_fim_isolado_espelha_para_ano_inicio():
    # Espelho de ano_inicio-só (já testado acima) -- informar só ano_fim
    # também preenche ano_inicio com o mesmo valor (período de 1 ano).
    filtro = FiltroRelatorio(escopo="brasil", ano_fim=2024)
    assert filtro.ano_inicio == 2024
    assert filtro.ano_fim == 2024


def test_formatar_ultima_acao_variacoes():
    concluida = AcaoMonitoramento(descricao="Vistoria", data_conclusao=date(2024, 3, 10))
    assert "concluída em" in _formatar_ultima_acao(concluida)

    prevista = AcaoMonitoramento(descricao="Vistoria", data_prevista=date(2024, 6, 1))
    assert "prevista para" in _formatar_ultima_acao(prevista)

    sem_data = AcaoMonitoramento(descricao="Vistoria")
    assert _formatar_ultima_acao(sem_data) == "Vistoria"


def test_formatar_ultimo_evento_com_e_sem_marco_catalogado():
    marco = MarcoCatalogo(id=1, codigo="fase_contratado", rotulo="Fase contratado")
    evento = EventoMarco(marco_id=1, data_ocorrencia=date(2024, 1, 5))
    assert "Fase contratado" in _formatar_ultimo_evento(evento, {1: marco})

    evento_sem_marco = EventoMarco(marco_id=999, data_prevista=date(2024, 2, 1))
    assert _formatar_ultimo_evento(evento_sem_marco, {}).startswith("Marco")


def test_data_inauguracao_previsao_variacoes():
    assert _data_inauguracao_previsao("Em andamento", []) == "—"

    concluido = EventoMarco(data_ocorrencia=date(2024, 5, 1))
    assert _data_inauguracao_previsao("Concluído", [concluido]) == formatar_data(concluido.data_ocorrencia)

    previsto = EventoMarco(data_prevista=date(2024, 7, 1))
    assert "previsão" in _data_inauguracao_previsao("Em andamento", [previsto])

    so_ocorrencia_sem_concluido = EventoMarco(data_ocorrencia=date(2024, 8, 1))
    assert _data_inauguracao_previsao("Em andamento", [so_ocorrencia_sem_concluido]) != "—"

    sem_datas = EventoMarco()
    assert _data_inauguracao_previsao("Em andamento", [sem_datas]) == "—"


def test_filtro_periodo_referencia_variacoes():
    assert FiltroRelatorio(escopo="brasil").periodo_referencia() is None
    assert FiltroRelatorio(escopo="brasil", ano_inicio=2023).periodo_referencia() == "2023"
    assert FiltroRelatorio(escopo="brasil", ano_inicio=2020, ano_fim=2023).periodo_referencia() == "2020 a 2023"


def test_assunto_brasil_e_valor_global_inconsistente():
    """Ramos de apresentação sem dependência de banco: Brasil não consulta
    estabelecimento, e o valor global menor que o repasse nunca é exibido."""
    assert _assunto(None, FiltroRelatorio(escopo="brasil")) == "Brasil"

    convenio = Convenio(valor_global=100, valor_repasse=200)
    assert _valor_global_confiavel(convenio) is None


def test_linhas_e_resumos_do_relatorio_preservam_detalhes_e_ausencias():
    convenio = Convenio(
        numero="123",
        convenente_nome="HOSPITAL TESTE",
        municipio="SAO PAULO",
        uf="SP",
        valor_global=200,
        valor_repasse=100,
    )
    assert len(_linhas_convenio(convenio, completo=False)) == 8
    linha_convenio_completa = _linhas_convenio(convenio, completo=True)
    assert linha_convenio_completa[0] == "123"
    assert "Sao Paulo" in linha_convenio_completa

    proposta = PropostaCandidata(
        id_proposta=456,
        nm_proponente="PROPONENTE TESTE",
        municipio="BRASILIA",
        uf="DF",
        tem_parceria=False,
    )
    assert len(_linhas_proposta(proposta, completo=False)) == 6
    linha_proposta_completa = _linhas_proposta(proposta, completo=True)
    assert linha_proposta_completa[0] == 456
    assert linha_proposta_completa[10] == "Não"

    assert _observacao_publica(None) is None
    assert _observacao_publica("Nota. Importado de Controle.xlsx sha256:abcdef, linha 24.") == "Nota."

    colunas, linhas = _resumo_equipamentos({1: [("Tomógrafo", 100.0), ("Tomógrafo", None), ("Ressonância", None)]})
    assert colunas == ["Equipamentos", "Quantidade", "Valor"]
    assert linhas == [["Tomógrafo", 2, 100.0], ["Ressonância", 1, None]]


@pytest.mark.db
def test_instrumentos_repasse_xlsx_simplificado_tem_resumo_convenios_propostas_monitoramento():
    db = SessionLocal()
    try:
        conteudo = montar_relatorio(
            db=db,
            formato="xlsx",
            nivel="simplificado",
            tipo_relatorio="instrumentos_repasse",
            filtro=FiltroRelatorio(escopo="brasil"),
        )
        pasta = load_workbook(BytesIO(conteudo))
        assert pasta.sheetnames == [
            "Resumo",
            "Resumo (Propostas)",
            "Resumo (Equipamentos)",
            "Convênios",
            "Propostas candidatas",
            "Monitoramento",
        ]
    finally:
        db.close()


@pytest.mark.db
def test_instrumentos_repasse_xlsx_completo_usa_timeline_de_monitoramento():
    db = SessionLocal()
    try:
        conteudo = montar_relatorio(
            db=db,
            formato="xlsx",
            nivel="completo",
            tipo_relatorio="instrumentos_repasse",
            filtro=FiltroRelatorio(escopo="uf", uf="DF"),
        )
        pasta = load_workbook(BytesIO(conteudo))
        assert "Monitoramento (timeline)" in pasta.sheetnames
        aba = pasta["Monitoramento (timeline)"]
        assert [c.value for c in aba[1]][:3] == ["Convênio", "Convenente", "Componente"]
    finally:
        db.close()


@pytest.mark.db
def test_instrumentos_repasse_docx_completo_narra_um_bloco_por_convenio():
    db = SessionLocal()
    try:
        conteudo = montar_relatorio(
            db=db,
            formato="docx",
            nivel="completo",
            tipo_relatorio="instrumentos_repasse",
            filtro=FiltroRelatorio(escopo="cnes", cnes="9000001"),
        )
        documento = Document(BytesIO(conteudo))
        textos = [p.text for p in documento.paragraphs]
        # Rótulo usa o `tipo_contratacao` real do convênio (Bloco 8 --
        # "PERSUS, TED e FAF não são tratados como convênio"), não mais
        # "Convênio" fixo -- fixture (`tests/fixtures_cobertura.py`) marca
        # este convênio como FAF.
        assert any("FAF __pytest_cnes_fk__" in t for t in textos)
        assert any(t.startswith("Objeto:") for t in textos)
        assert any(t == "Financeiro" for t in textos)
        # Algumas tabelas do documento têm só cabeçalho (ex. "Resumo
        # (Equipamentos)" quando o convênio da fixture não tem
        # EquipamentoMarcador) -- `rows[1]` estouraria IndexError nelas, daí
        # o filtro por `len(tabela.rows) > 1` antes de olhar o conteúdo.
        assert any(
            "Valor Global" in "".join(c.text for c in tabela.rows[1].cells)
            for tabela in documento.tables
            if len(tabela.rows) > 1
        )
    finally:
        db.close()


@pytest.mark.db
def test_instrumentos_repasse_docx_simplificado_mostra_estado_atual_sem_timeline():
    """Fase 2 (Plan Mode relatorios) -- nível Simplificado continua mostrando
    só o estado atual (4 rótulos), nunca a timeline inteira. Instrumento
    948686 (`scripts/seed_monitoramento.py`) tem 3 eventos reais."""
    db = SessionLocal()
    try:
        _garantir_convenio_948686(db)
        conteudo = montar_relatorio(
            db=db,
            formato="docx",
            nivel="simplificado",
            tipo_relatorio="instrumentos_repasse",
            filtro=FiltroRelatorio(escopo="uf", uf="DF"),
        )
        documento = Document(BytesIO(conteudo))
        textos = [p.text for p in documento.paragraphs]
        assert any(t.startswith("Último Evento:") for t in textos)
        assert any(t.startswith("Última Ação:") for t in textos)
        celulas_tabelas = ["".join(c.text for c in row.cells) for tabela in documento.tables for row in tabela.rows]
        assert not any("Matrícula CNEN" in linha for linha in celulas_tabelas)
        db.rollback()  # nunca commita dado de teste
    finally:
        db.close()


@pytest.mark.db
def test_instrumentos_repasse_docx_completo_mostra_timeline_inteira_e_acoes():
    """Fase 2 -- nível Completo troca as 4 linhas de estado atual pela
    timeline INTEIRA de eventos (não só o último) mais a lista completa de
    ações, "mais etapas, não só a última" (pedido do usuário)."""
    db = SessionLocal()
    try:
        _garantir_convenio_948686(db)
        instrumento = db.query(InstrumentoEquipamento).filter_by(nr_convenio="948686").one()
        acao = AcaoMonitoramento(
            instrumento_id=instrumento.id,
            descricao="Cobrar retorno da CNEN sobre o RPAS",
            data_prevista=date(2026, 10, 1),
            responsavel="PRISCILA",
        )
        db.add(acao)
        db.flush()

        conteudo = montar_relatorio(
            db=db,
            formato="docx",
            nivel="completo",
            tipo_relatorio="instrumentos_repasse",
            filtro=FiltroRelatorio(escopo="uf", uf="DF"),
        )
        documento = Document(BytesIO(conteudo))
        textos = [p.text for p in documento.paragraphs]

        # Não mostra mais só "a última" -- as 2 etiquetas de estado atual
        # somem no nível Completo (viram tabela).
        assert not any(t.startswith("Último Evento:") for t in textos)
        assert not any(t.startswith("Última Ação:") for t in textos)

        # Timeline inteira: os 3 eventos reais do seed aparecem como linhas
        # de tabela (marco "Matrícula CNEN" é só 1 dos 3, prova que não é
        # só "o mais recente").
        celulas_tabelas = ["".join(c.text for c in row.cells) for tabela in documento.tables for row in tabela.rows]
        assert any("Matrícula CNEN" in linha for linha in celulas_tabelas)
        assert any("SCRA modificação/casamata" in linha for linha in celulas_tabelas)

        # Lista completa de ações -- a ação recém-criada aparece com
        # descrição e responsável.
        assert any("Cobrar retorno da CNEN sobre o RPAS" in linha for linha in celulas_tabelas)
        assert any("PRISCILA" in linha for linha in celulas_tabelas)
        db.rollback()  # nunca commita dado de teste
    finally:
        db.close()


@pytest.mark.db
def test_analise_merito_xlsx_so_tem_abas_de_cobertura():
    db = SessionLocal()
    try:
        conteudo = montar_relatorio(
            db=db,
            formato="xlsx",
            nivel="simplificado",
            tipo_relatorio="analise_merito",
            filtro=FiltroRelatorio(escopo="uf", uf="DF"),
        )
        pasta = load_workbook(BytesIO(conteudo))
        assert all(nome.startswith("Cobertura") for nome in pasta.sheetnames)
        assert "Convênios" not in pasta.sheetnames
        assert "Monitoramento" not in pasta.sheetnames
    finally:
        db.close()


@pytest.mark.db
def test_analise_merito_escopo_cnes_e_rejeitado():
    db = SessionLocal()
    try:
        with pytest.raises(ValidationError):
            montar_relatorio(
                db=db,
                formato="xlsx",
                nivel="simplificado",
                tipo_relatorio="analise_merito",
                filtro=FiltroRelatorio(escopo="cnes", cnes="9000001"),
            )
    finally:
        db.close()


@pytest.mark.db
def test_analise_merito_uf_sem_cobertura_gera_arquivo_valido_sem_quebrar():
    db = SessionLocal()
    try:
        conteudo = montar_relatorio(
            db=db,
            formato="xlsx",
            nivel="simplificado",
            tipo_relatorio="analise_merito",
            # UF sem nenhum dado de cobertura no seed sintético (só DF tem
            # linha de MacroCoverage) -- a família TOMOGRAFO existe, só não
            # tem dado pro AC; devolve aba com só cabeçalho, sem quebrar.
            filtro=FiltroRelatorio(escopo="uf", uf="AC"),
        )
        pasta = load_workbook(BytesIO(conteudo))
        assert pasta.sheetnames == ["Cobertura TOMOGRAFO"]
        assert pasta["Cobertura TOMOGRAFO"].max_row == 1
    finally:
        db.close()


@pytest.mark.db
def test_analise_merito_docx_escopo_regiao_nivel_completo():
    # Achado real: nenhum teste exercitava o formato docx de analise_merito
    # nem o escopo=regiao/nivel=completo (granularidade de município).
    db = SessionLocal()
    try:
        conteudo = montar_relatorio(
            db=db,
            formato="docx",
            nivel="completo",
            tipo_relatorio="analise_merito",
            filtro=FiltroRelatorio(escopo="regiao", regiao="Centro-Oeste"),
        )
        documento = Document(BytesIO(conteudo))
        textos = " ".join(p.text for p in documento.paragraphs)
        assert "Análise de mérito" in textos
        assert "Região Centro-Oeste" in textos
    finally:
        db.close()


@pytest.mark.db
def test_instrumentos_repasse_docx_completo_narra_bloco_de_proposta():
    # Achado real: nenhum teste exercitava `_bloco_proposta_narrativo` --
    # o docx completo só era testado com escopo=cnes, que nunca tem
    # proposta candidata (PropostaCandidata não tem granularidade de CNES
    # própria na maioria dos casos).
    db = SessionLocal()
    proposta = None
    try:
        proposta = PropostaCandidata(
            id_proposta=int(str(uuid4().int)[:9]),
            cnpj_ente_recebedor="00000000000100",
            nm_proponente="Proponente Pytest Narrativa — apagar",
            municipio="Brasília",
            uf="DF",
            ds_objeto="Objeto de teste",
            nm_programa="Programa de teste",
            id_programa=1,
            componente_batido="Tomógrafo",
            vl_global_proposta=1000,
            situacao_proposta="Em análise",
            tem_parceria=False,
        )
        db.add(proposta)
        db.commit()

        conteudo = montar_relatorio(
            db=db,
            formato="docx",
            nivel="completo",
            tipo_relatorio="instrumentos_repasse",
            filtro=FiltroRelatorio(escopo="uf", uf="DF"),
        )
        documento = Document(BytesIO(conteudo))
        textos = [p.text for p in documento.paragraphs]
        assert any("Proponente Pytest Narrativa" in t for t in textos)
        assert any(t == "Parceria formalizada: Não" for t in textos)
    finally:
        if proposta is not None:
            db.query(PropostaCandidata).filter_by(id=proposta.id).delete()
            db.commit()
        db.close()


@pytest.mark.db
def test_instrumentos_repasse_xlsx_escopo_municipio():
    # Achado real: nenhum teste exercitava escopo=municipio via
    # montar_relatorio -- só a validação isolada de FiltroRelatorio.
    db = SessionLocal()
    try:
        conteudo = montar_relatorio(
            db=db,
            formato="xlsx",
            nivel="simplificado",
            tipo_relatorio="instrumentos_repasse",
            filtro=FiltroRelatorio(escopo="municipio", municipio="Brasília", uf="DF"),
        )
        pasta = load_workbook(BytesIO(conteudo))
        assert "Resumo" in pasta.sheetnames
        assert "Convênios" in pasta.sheetnames
    finally:
        db.close()


@pytest.mark.db
def test_instrumentos_repasse_docx_assunto_varia_por_escopo():
    # Achado real: `_assunto` (cabeçalho "Assunto:" do Word) só era
    # exercitado pro escopo=cnes -- município/UF/região nunca passavam
    # pelo docx (só testados via xlsx ou via FiltroRelatorio isolado).
    db = SessionLocal()
    try:
        for filtro, esperado in (
            (FiltroRelatorio(escopo="municipio", municipio="Brasília", uf="DF"), "Brasília/DF"),
            (FiltroRelatorio(escopo="uf", uf="DF"), "UF DF"),
            (FiltroRelatorio(escopo="regiao", regiao="Centro-Oeste"), "Região Centro-Oeste"),
        ):
            conteudo = montar_relatorio(
                db=db, formato="docx", nivel="simplificado", tipo_relatorio="instrumentos_repasse", filtro=filtro
            )
            textos = [p.text for p in Document(BytesIO(conteudo)).paragraphs]
            assert any(f"Assunto: {esperado}" in t for t in textos)
    finally:
        db.close()


@pytest.mark.db
def test_analise_merito_sem_familia_publicada_gera_arquivo_com_aviso_em_vez_de_quebrar(monkeypatch):
    import app.services.relatorios as relatorios_mod

    monkeypatch.setattr(relatorios_mod.execucoes_repo, "listar_familias_publicadas", lambda db: [])
    db = SessionLocal()
    try:
        conteudo = montar_relatorio(
            db=db,
            formato="xlsx",
            nivel="simplificado",
            tipo_relatorio="analise_merito",
            filtro=FiltroRelatorio(escopo="brasil"),
        )
        pasta = load_workbook(BytesIO(conteudo))
        assert pasta.sheetnames == ["Cobertura"]
        assert pasta["Cobertura"]["A2"].value.startswith("Nenhuma família")
    finally:
        db.close()


@pytest.mark.db
def test_montar_relatorio_docx_titulo_do_filtro():
    db = SessionLocal()
    try:
        conteudo = montar_relatorio(
            db=db,
            formato="docx",
            nivel="completo",
            tipo_relatorio="instrumentos_repasse",
            filtro=FiltroRelatorio(escopo="uf", uf="DF"),
        )
        documento = Document(BytesIO(conteudo))
        assert documento.paragraphs[0].text == "Instrumentos e repasse -- UF DF"
    finally:
        db.close()
