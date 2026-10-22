"""Paridade de conteúdo/layout e liberação de XML na gravação progressiva."""

from io import BytesIO
from zipfile import ZipFile

import pytest
from docx import Document
from lxml import etree

from app.reports import docx_builder
from app.reports.docx_tabelado import SecaoTabela, gerar_bytes_tabelados


def estrutura(elemento):
    return (elemento.tag, sorted(elemento.attrib.items()), elemento.text, [estrutura(c) for c in elemento])


@pytest.mark.parametrize("total", [0, 1, 4, 5, 6, 65, 257])
@pytest.mark.parametrize("formatado", [False, True])
def test_gravacao_progressiva_preserva_xml_e_partes_do_modelo(total, formatado):
    # Arrange: texto com acentos, caracteres XML, quebras e identificadores.
    linhas = [[f"Município <{n}> & ação\nsegunda linha", "0000434", n, None] for n in range(total)]
    secao = SecaoTabela(
        "Cobertura",
        ["Nome", "CNES", "Quantidade", "Ausente"],
        linhas,
        [3, 1, 1, 1] if formatado else None,
        11 if formatado else None,
    )
    comum = docx_builder.novo_documento("Relatório")
    docx_builder.adicionar_titulo(comum, secao.titulo)
    docx_builder.adicionar_tabela(
        comum, secao.colunas, linhas, pesos_colunas=secao.pesos_colunas, fonte_pt=secao.fonte_pt
    )
    progressivo = docx_builder.novo_documento("Relatório")
    # Act
    antigo = docx_builder.gerar_bytes(comum)
    novo = gerar_bytes_tabelados(progressivo, [secao])
    # Assert: comparação XML estrutural, incluindo estilos, larguras e paginação.
    with ZipFile(BytesIO(antigo)) as a, ZipFile(BytesIO(novo)) as b:
        assert b.testzip() is None
        assert a.namelist() == b.namelist()
        assert estrutura(etree.fromstring(a.read("word/document.xml"))) == estrutura(
            etree.fromstring(b.read("word/document.xml"))
        )
        for nome in a.namelist():
            if nome != "word/document.xml":
                assert a.read(nome) == b.read(nome)
    reaberto = Document(BytesIO(novo))
    assert len(reaberto.tables[0].rows) == total + 1
    assert len(progressivo.tables[0].rows) == 1  # XML dos dados não fica no template.
    if total:
        assert reaberto.tables[0].cell(1, 1).text == "0000434"


def test_multiplas_secoes_preservam_ordem_e_nao_misturam_linhas():
    # Arrange
    secoes = [
        SecaoTabela("Primeira", ["Nome"], [["ação"]]),
        SecaoTabela("Vazia", ["Nome"], []),
        SecaoTabela("Última", ["Nome"], [["conclusão"]]),
    ]
    # Act
    arquivo = gerar_bytes_tabelados(docx_builder.novo_documento("Capa"), secoes)
    # Assert
    reaberto = Document(BytesIO(arquivo))
    assert [p.text for p in reaberto.paragraphs] == ["Capa", "Primeira", "Vazia", "Última"]
    assert [[c.text for r in t.rows for c in r.cells] for t in reaberto.tables] == [
        ["Nome", "ação"],
        ["Nome"],
        ["Nome", "conclusão"],
    ]


def test_linha_invalida_nao_impede_proxima_geracao():
    # Arrange
    documento = docx_builder.novo_documento("Entrada inválida")
    # Act / Assert: zip strict rejeita formato inválido, sem devolver pacote parcial.
    with pytest.raises(ValueError):
        gerar_bytes_tabelados(documento, [SecaoTabela("Tabela", ["Nome"], [["um", "dois"]])])
    assert len(documento.tables[0].rows) == 1
    valido = gerar_bytes_tabelados(docx_builder.novo_documento("Válido"), [SecaoTabela("Tabela", ["Nome"], [["ok"]])])
    assert Document(BytesIO(valido)).tables[0].cell(1, 0).text == "ok"


def test_recusa_template_com_tabela_existente_sem_perder_conteudo():
    # Arrange
    documento = docx_builder.novo_documento("Já preenchido")
    docx_builder.adicionar_tabela(documento, ["Nome"], [["preservar"]])
    # Act / Assert
    with pytest.raises(ValueError, match="iniciar sem tabelas"):
        gerar_bytes_tabelados(documento, [])
    assert documento.tables[0].cell(1, 0).text == "preservar"
