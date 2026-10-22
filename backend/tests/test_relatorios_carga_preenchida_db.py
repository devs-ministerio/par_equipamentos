"""Teto sintético com histórico; RSS é pico do processo, não custo isolado."""

import json
import resource
import sys
from io import BytesIO
from time import perf_counter

import pytest
from docx import Document
from openpyxl import load_workbook

from app.db.base import SessionLocal
from app.services.relatorios import FiltroRelatorio, montar_relatorio
from tests.fixtures_relatorios_preenchidos import criar_carteira_preenchida
from tests.numeros_relatorio import numeros_dos_blocos
from tests.test_relatorios_lotes_db import contar_selects


@pytest.mark.parametrize("formato", ["docx", "xlsx"])
def test_carteira_preenchida_de_500_preserva_todos_os_instrumentos(formato, tmp_path):
    # Given: 1.000 itens, 1.500 eventos ativos, 1.000 ações ativas e histórico.
    with SessionLocal() as db:
        try:
            criar_carteira_preenchida(db, 500)
            # When
            inicio = perf_counter()
            with contar_selects(db) as consultas:
                conteudo = montar_relatorio(
                    db=db,
                    formato=formato,
                    nivel="completo",
                    tipo_relatorio="instrumentos_repasse",
                    filtro=FiltroRelatorio(escopo="uf", uf="AC"),
                )
            duracao = perf_counter() - inicio
            rss = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss
            (tmp_path / "carga-preenchida.json").write_text(
                json.dumps(
                    {
                        "formato": formato,
                        "instrumentos": 500,
                        "geracao_segundos": duracao,
                        "selects": len(consultas),
                        "bytes": len(conteudo),
                        "rss_pico_processo_bytes": rss if sys.platform == "darwin" else rss * 1024,
                    }
                ),
                encoding="utf-8",
            )
            # Then: invariantes sobre todos os registros, sem um prazo flaky.
            esperados = {f"QA_CARGA_{n:04}" for n in range(500)}
            if formato == "docx":
                documento = Document(BytesIO(conteudo))
                numeros = numeros_dos_blocos(documento.paragraphs)
                assert len(numeros) == 500 and set(numeros) == esperados
                celulas = [c.text for t in documento.tables for r in t.rows for c in r.cells]
                for n in range(500):
                    assert f"Concluída QA {n:04}" in celulas
                    assert f"Pendente QA {n:04}" in celulas
                assert not any("HISTORICO_" in c for c in celulas)
                itens = [
                    tuple(c.text for c in row.cells)
                    for tabela in documento.tables
                    if tabela.rows[0].cells[0].text == "Equipamentos"
                    for row in list(tabela.rows)[1:]
                ]
                # O resumo executivo usa o mesmo cabeçalho dos 500 blocos.
                # Conferir tanto as duas linhas agregadas quanto as 1.000 individuais.
                assert len(itens) == 1002
                assert itens.count(("Equipamento QA carga", "1", "R$ 0,00")) == 500
                assert itens.count(("Equipamento QA auxiliar", "1", "—")) == 500
                assert itens.count(("Equipamento QA carga", "500", "R$ 0,00")) == 1
                assert itens.count(("Equipamento QA auxiliar", "500", "—")) == 1
                assert len(consultas) == 12
            else:
                pasta = load_workbook(BytesIO(conteudo))
                try:
                    linhas = list(pasta["Monitoramento (timeline)"].iter_rows(min_row=2, values_only=True))
                    assert len(linhas) == 1500
                    assert {r[0] for r in linhas} == esperados
                    assert not any("HISTORICO_" in str(r) for r in linhas)
                    equipamentos = set(pasta["Resumo (Equipamentos)"].iter_rows(min_row=2, values_only=True))
                    assert equipamentos == {
                        ("Equipamento QA carga", 500, 0),
                        ("Equipamento QA auxiliar", 500, None),
                    }
                    assert len(consultas) == 9
                finally:
                    pasta.close()
        finally:
            db.rollback()
