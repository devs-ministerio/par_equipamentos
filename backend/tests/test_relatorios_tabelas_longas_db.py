"""Amostras para revisão visual, com asserções sobre os arquivos reais."""

from io import BytesIO

import pytest
from docx import Document
from openpyxl import load_workbook

from app.db.base import SessionLocal
from app.services.relatorios import FiltroRelatorio, montar_relatorio
from tests.fixtures_relatorios_longos import criar_monitoramento_longo


@pytest.mark.parametrize("formato", ["docx", "xlsx"])
def test_tabelas_longas_preservam_catalogo_e_dados_textuais(formato, tmp_path):
    # Given: catálogo fixo, marcos distintos, identificadores e tarefas sintéticos.
    with SessionLocal() as db:
        marcos, acoes = criar_monitoramento_longo(db)
        try:
            # When
            conteudo = montar_relatorio(
                db=db,
                formato=formato,
                nivel="completo",
                tipo_relatorio="instrumentos_repasse",
                filtro=FiltroRelatorio(escopo="cnes", cnes="0000788"),
            )
            # Then: conteúdo completo no arquivo, não apenas arquivo não vazio.
            if formato == "docx":
                documento = Document(BytesIO(conteudo))
                celulas = [cell.text for table in documento.tables for row in table.rows for cell in row.cells]
                assert set(marcos).issubset({c.split("\n", 1)[0] for c in celulas})
                assert any(c.endswith("\nRegulatório") for c in celulas)
                assert "02/01/2026" in celulas
                assert "31/12/2027" in celulas
                assert set(acoes).issubset(celulas)
                for tabela in documento.tables:
                    assert tabela.rows[0]._tr.xpath("./w:trPr/w:tblHeader")
                    assert all(row._tr.xpath("./w:trPr/w:cantSplit") for row in tabela.rows)
            else:
                pasta = load_workbook(BytesIO(conteudo))
                aba = pasta["Monitoramento (timeline)"]
                linhas = list(aba.iter_rows(min_row=2, values_only=True))
                assert {linha[3] for linha in linhas} == set(marcos)
                assert len(linhas) == len(marcos)
                assert {linha[0] for linha in linhas} == {"QA_LONGO_001"}
                assert aba.freeze_panes == "A2"
            (tmp_path / f"monitoramento-longo.{formato}").write_bytes(conteudo)
        finally:
            db.rollback()
