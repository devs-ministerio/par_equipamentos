"""Recortes e conteúdo de propostas/monitoramento nos relatórios finais."""

from datetime import date
from io import BytesIO

import pytest
from docx import Document
from openpyxl import load_workbook

from app.db.base import SessionLocal
from app.db.models import CnesEstabelecimento, Convenio, InstrumentoEquipamento
from app.repositories.convenios import listar_numeros_fora_escopo
from app.services import monitoramento_instrumentos
from app.services.relatorios import FiltroRelatorio, montar_relatorio
from tests.fixtures_relatorios_cruzados import criar_carteira_cruzada

CASOS = [
    ("brasil", FiltroRelatorio(escopo="brasil"), (1, 2, 3, 4, 5), (1, 2, 3, 4, 5)),
    ("nordeste", FiltroRelatorio(escopo="regiao", regiao="Nordeste"), (1, 2, 3, 5), (1, 2, 3, 5)),
    ("uf", FiltroRelatorio(escopo="uf", uf="BA"), (1, 3, 5), (1, 3, 5)),
    ("municipio", FiltroRelatorio(escopo="municipio", uf="SP", municipio="São Paulo"), (4,), (4,)),
    ("cnes", FiltroRelatorio(escopo="cnes", cnes="0000701"), (1, 3, 5), (1, 3, 5)),
    ("ano_origem", FiltroRelatorio(escopo="brasil", ano_inicio=2025), (1, 4), (1, 4)),
    ("historico", FiltroRelatorio(escopo="brasil", ano_fim=2024), (2,), (2,)),
    ("ano_corrente", FiltroRelatorio(escopo="brasil", ano_inicio=2026), (3,), (1, 3, 4, 5)),
    ("periodo", FiltroRelatorio(escopo="brasil", ano_inicio=2024, ano_fim=2026), (1, 2, 3, 4), (1, 2, 3, 4, 5)),
    ("vazio", FiltroRelatorio(escopo="cnes", cnes="9999999"), (), ()),
]


@pytest.fixture
def carteira_cruzada(monkeypatch):
    class RelogioSintetico(date):
        @classmethod
        def today(cls):
            return cls(2026, 10, 19)

    # Somente relógio é substituído; SQL/services/builders rodam reais.
    monkeypatch.setattr(monitoramento_instrumentos, "date", RelogioSintetico)
    with SessionLocal() as db:
        criar_carteira_cruzada(db)
        try:
            yield db
        finally:
            db.rollback()


def numeros_monitorados_word(documento):
    numero = None
    numeros = set()
    aguardando_fase = False
    for paragrafo in documento.paragraphs:
        if paragrafo.text.startswith("Número: "):
            numero = paragrafo.text.removeprefix("Número: ")
        if aguardando_fase:
            if paragrafo.text.startswith("Situação: ") and paragrafo.text != "Situação: —":
                numeros.add(numero)
            aguardando_fase = False
        if paragrafo.text == "Monitoramento Interno":
            aguardando_fase = True
    return numeros


@pytest.mark.parametrize("formato", ["xlsx", "docx"])
@pytest.mark.parametrize("nivel", ["simplificado", "completo"])
@pytest.mark.parametrize("nome,filtro,propostas,monitorados", CASOS, ids=[c[0] for c in CASOS])
def test_recorte_preserva_propostas_e_monitorados_sem_reintroduzir_obras(
    carteira_cruzada,
    nome,
    filtro,
    propostas,
    monitorados,
    nivel,
    formato,
):
    # Given: identidades da fixture distinguem os demais seeds válidos do banco.
    universo_monitorados = {f"900700{n}" for n in range(1, 7)}
    universo_propostas = {str(9107000 + n) for n in range(1, 6)}
    # When
    conteudo = montar_relatorio(
        db=carteira_cruzada, formato=formato, nivel=nivel, tipo_relatorio="instrumentos_repasse", filtro=filtro
    )
    # Then: oráculos explícitos, incluindo a exceção anual do acompanhamento em aberto.
    if formato == "xlsx":
        pasta = load_workbook(BytesIO(conteudo))
        propostas_obtidas = {str(row[0].value) for row in pasta["Propostas candidatas"].iter_rows(min_row=2)}
        aba = pasta["Monitoramento" if nivel == "simplificado" else "Monitoramento (timeline)"]
        mon_obtidos = {str(row[0].value) for row in aba.iter_rows(min_row=2)}
    else:
        documento = Document(BytesIO(conteudo))
        propostas_obtidas = {
            p.text.rsplit(" -- Proposta ", 1)[1] for p in documento.paragraphs if " -- Proposta " in p.text
        }
        mon_obtidos = numeros_monitorados_word(documento)
    assert propostas_obtidas & universo_propostas == {str(9107000 + n) for n in propostas}
    assert mon_obtidos & universo_monitorados == {f"900700{n}" for n in monitorados}


@pytest.mark.parametrize("nivel", ["simplificado", "completo"])
def test_word_nao_omite_monitorado_sem_convenio_nem_inventa_financeiro(carteira_cruzada, nivel, tmp_path):
    # Given: o cadastro de monitoramento é independente da existência de Convênio.
    assert carteira_cruzada.query(Convenio).filter_by(numero="9007999").count() == 0
    carteira_cruzada.add(CnesEstabelecimento(cnes="0000799", nome_estabelecimento="Unidade sem espelho"))
    carteira_cruzada.flush()
    carteira_cruzada.add(
        InstrumentoEquipamento(
            nr_convenio="9007999",
            nome_convenente="Instituição sintética sem espelho",
            cnes="0000799",
            municipio="SALVADOR",
            uf="BA",
            tipo_contratacao="FAF",
        )
    )
    carteira_cruzada.flush()
    # When
    conteudo = montar_relatorio(
        db=carteira_cruzada,
        formato="docx",
        nivel=nivel,
        tipo_relatorio="instrumentos_repasse",
        filtro=FiltroRelatorio(escopo="cnes", cnes="0000799"),
    )
    documento = Document(BytesIO(conteudo))
    # Then
    assert numeros_monitorados_word(documento) == {"9007999"}
    assert any("Instituição sintética sem espelho" in p.text for p in documento.paragraphs)
    assert not any(p.text == "Financeiro" for p in documento.paragraphs)
    celulas = [cell.text for table in documento.tables for row in table.rows for cell in row.cells]
    assert "Valor global" not in celulas
    if nivel == "completo":
        assert "Marco" not in celulas  # Sem tabela de sete colunas vazia.
        assert any(p.text == "— nenhum evento registrado —" for p in documento.paragraphs)
        assert any(p.text == "— nenhuma ação registrada —" for p in documento.paragraphs)
    (tmp_path / f"monitoramento-sem-espelho-{nivel}.docx").write_bytes(conteudo)


def test_exclusao_em_lote_preserva_identidade_sem_espelho(carteira_cruzada):
    assert listar_numeros_fora_escopo(carteira_cruzada, set()) == set()
    assert listar_numeros_fora_escopo(carteira_cruzada, {"9007001", "9007006", "9007999"}) == {"9007006"}
    # Exportar não pode apagar ou alterar a decisão de escopo nem o cadastro interno.
    montar_relatorio(
        db=carteira_cruzada,
        formato="xlsx",
        nivel="completo",
        tipo_relatorio="instrumentos_repasse",
        filtro=FiltroRelatorio(escopo="brasil"),
    )
    assert carteira_cruzada.query(Convenio).filter_by(numero="9007006", fora_escopo_equipamentos=True).count() == 1
    assert carteira_cruzada.query(InstrumentoEquipamento).filter_by(nr_convenio="9007006").count() == 1


@pytest.mark.parametrize("formato", ["xlsx", "docx"])
def test_acompanhamento_corrente_muda_na_virada_do_ano(carteira_cruzada, monkeypatch, formato):
    class AnoSeguinte(date):
        @classmethod
        def today(cls):
            return cls(2027, 1, 1)

    monkeypatch.setattr(monitoramento_instrumentos, "date", AnoSeguinte)
    conteudo = montar_relatorio(
        db=carteira_cruzada,
        formato=formato,
        nivel="completo",
        tipo_relatorio="instrumentos_repasse",
        filtro=FiltroRelatorio(escopo="cnes", cnes="0000701", ano_inicio=2027),
    )
    if formato == "xlsx":
        pasta = load_workbook(BytesIO(conteudo))
        obtidos = {str(row[0].value) for row in pasta["Monitoramento (timeline)"].iter_rows(min_row=2)}
        assert pasta["Propostas candidatas"].max_row == 1
    else:
        documento = Document(BytesIO(conteudo))
        obtidos = numeros_monitorados_word(documento)
        assert not any(" -- Proposta " in p.text for p in documento.paragraphs)
    assert obtidos == {"9007001", "9007005"}
