"""Carga preenchida e número de consultas real, sem mocks de domínio."""

import json
from contextlib import contextmanager
from io import BytesIO

import pytest
from docx import Document
from openpyxl import load_workbook
from sqlalchemy import event

from app.db.base import SessionLocal
from app.repositories import monitoramento as repo
from app.services.relatorios import FiltroRelatorio, montar_relatorio
from app.services.relatorios_monitoramento import carregar_monitoramento_relatorio
from tests.fixtures_relatorios_preenchidos import criar_carteira_preenchida


@contextmanager
def contar_selects(db):
    consultas = []

    def capturar(_conn, _cursor, statement, _parameters, _context, _executemany):
        if statement.lstrip().upper().startswith("SELECT"):
            consultas.append(1)

    conexao = db.connection()
    event.listen(conexao, "before_cursor_execute", capturar)
    try:
        yield consultas
    finally:
        event.remove(conexao, "before_cursor_execute", capturar)


@pytest.mark.parametrize("formato,nivel", [("docx", "simplificado"), ("docx", "completo"), ("xlsx", "completo")])
def test_consultas_do_relatorio_nao_crescem_por_instrumento(formato, nivel, tmp_path):
    # Given
    with SessionLocal() as db:
        try:
            criar_carteira_preenchida(db, 8)
            contagens = []
            arquivos = []
            # When: um mesmo recorte geográfico; apenas a cardinalidade difere.
            for filtro in (FiltroRelatorio(escopo="cnes", cnes="0000810"), FiltroRelatorio(escopo="uf", uf="AC")):
                with contar_selects(db) as consultas:
                    arquivos.append(
                        montar_relatorio(
                            db=db, formato=formato, nivel=nivel, tipo_relatorio="instrumentos_repasse", filtro=filtro
                        )
                    )
                contagens.append(len(consultas))
            (tmp_path / "consultas.json").write_text(json.dumps(contagens), encoding="utf-8")
            # Then: conteúdo cresce, consultas não; não é threshold de duração.
            # CNES faz ainda uma leitura do nome da unidade no assunto DOCX.
            # Crescer a carteira não pode acrescentar consultas por instrumento.
            assert contagens[1] <= contagens[0], contagens
            if formato == "docx":
                documento = Document(BytesIO(arquivos[1]))
                textos = "\n".join(p.text for p in documento.paragraphs)
                celulas = [c.text for t in documento.tables for r in t.rows for c in r.cells]
                for n in range(8):
                    assert f"Número: QA_CARGA_{n:04}" in textos
                    if nivel == "completo":
                        assert f"Concluída QA {n:04}" in celulas
                        assert f"Pendente QA {n:04}" in celulas
                    else:
                        assert f"Última Ação: Concluída QA {n:04}" in textos
                if nivel == "simplificado":
                    assert textos.count("(03/01/2027)") == 8
                    assert textos.count("Data de Inauguração/Previsão: 03/01/2027 (previsão)") == 8
                if nivel == "completo":
                    assert "Responsável QA cadastrado" in celulas
                    assert "Responsável QA legado" in celulas
                    assert celulas.index("Concluída QA 0000") < celulas.index("Pendente QA 0000")
                assert "HISTORICO_" not in textos + "\n".join(celulas)
                assert "Equipamento QA carga" in celulas
                assert any(
                    [c.text for c in row.cells] == ["Equipamento QA carga", "1", "R$ 0,00"]
                    for tabela in documento.tables
                    for row in tabela.rows
                )
                assert any(
                    [c.text for c in row.cells] == ["Equipamento QA auxiliar", "1", "—"]
                    for tabela in documento.tables
                    for row in tabela.rows
                )
            else:
                pasta = load_workbook(BytesIO(arquivos[1]))
                try:
                    linhas = list(pasta["Monitoramento (timeline)"].iter_rows(min_row=2, values_only=True))
                    assert len(linhas) == 24
                    assert {r[0] for r in linhas} == {f"QA_CARGA_{n:04}" for n in range(8)}
                    assert not any("HISTORICO_" in str(r) for r in linhas)
                finally:
                    pasta.close()
        finally:
            db.rollback()


@pytest.mark.parametrize("ativos", [True, False])
def test_lotes_preservam_ordenacao_historico_e_isolamento(ativos):
    # Given
    with SessionLocal() as db:
        try:
            instrumentos = criar_carteira_preenchida(db, 3)
            ids = [instrumentos[0].id, instrumentos[2].id]
            # When
            eventos = repo.listar_eventos_por_instrumentos(db, [*ids, ids[0]], apenas_ativos=ativos)
            acoes = repo.listar_acoes_por_instrumentos(db, ids, apenas_ativas=ativos)
            # Then: sem duplicação ou dados de outro instrumento.
            assert set(eventos) == set(acoes) == set(ids)
            for instrumento_id in ids:
                assert len(eventos[instrumento_id]) == (3 if ativos else 5)
                assert len(acoes[instrumento_id]) == (2 if ativos else 4)
                assert eventos[instrumento_id] == sorted(
                    eventos[instrumento_id], key=lambda e: (e.created_at, e.id), reverse=True
                )
                assert [a.descricao for a in acoes[instrumento_id]][0].startswith("Concluída QA")
                assert (
                    repo.listar_eventos_do_instrumento(db, instrumento_id, apenas_ativos=ativos)
                    == eventos[instrumento_id]
                )
                assert (
                    repo.listar_acoes_do_instrumento(db, instrumento_id, apenas_ativas=ativos) == acoes[instrumento_id]
                )
            assert repo.listar_eventos_por_instrumentos(db, [-1]) == {}
            assert repo.listar_acoes_por_instrumentos(db, [-1]) == {}
        finally:
            db.rollback()


def test_recorte_vazio_nao_consulta_catalogo_historico_ou_nomes():
    # Given
    with SessionLocal() as db:
        # When
        with contar_selects(db) as consultas:
            assert repo.listar_eventos_por_instrumentos(db, []) == {}
            assert repo.listar_acoes_por_instrumentos(db, []) == {}
            dados = carregar_monitoramento_relatorio(db=db, instrumento_ids=[])
        # Then
        assert consultas == []
        assert (
            dados.marcos
            == dados.eventos_por_instrumento
            == dados.acoes_por_instrumento
            == dados.nomes_responsaveis
            == {}
        )
