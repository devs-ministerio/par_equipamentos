"""Valores exportados de fato, em quatro versões e PostgreSQL real."""

from dataclasses import replace
from datetime import datetime
from io import BytesIO

import pytest
from docx import Document
from openpyxl import load_workbook

from app.db.base import SessionLocal
from app.db.models import Convenio
from app.services.relatorios import FiltroRelatorio, montar_relatorio
from tests.fixtures_relatorios_matriz import EQUIPAMENTO, PREFIXO, criar_carteira_matriz

CASOS = [
    ("brasil", FiltroRelatorio(escopo="brasil"), (1, 2, 3, 4, 5)),
    ("regiao", FiltroRelatorio(escopo="regiao", regiao="Nordeste"), (1, 2, 3, 5)),
    ("uf", FiltroRelatorio(escopo="uf", uf="BA"), (1, 3, 5)),
    ("municipio", FiltroRelatorio(escopo="municipio", uf="BA", municipio="salvador"), (1, 3, 5)),
    ("municipio_acento", FiltroRelatorio(escopo="municipio", uf="SP", municipio="São Paulo"), (4,)),
    ("cnes", FiltroRelatorio(escopo="cnes", cnes="0000701"), (1, 3, 5)),
    ("inicio_isolado", FiltroRelatorio(escopo="brasil", ano_inicio=2025), (1, 4)),
    ("fim_isolado", FiltroRelatorio(escopo="brasil", ano_fim=2024), (2,)),
    ("periodo_inclusivo", FiltroRelatorio(escopo="brasil", ano_inicio=2024, ano_fim=2025), (1, 2, 4)),
    ("situacao", FiltroRelatorio(escopo="brasil", situacao="Concluído"), (3,)),
    ("programa", FiltroRelatorio(escopo="brasil", programa="Programa matriz A"), (1, 2, 5)),
    ("contratacao", FiltroRelatorio(escopo="brasil", tipo_contratacao="TED"), (3,)),
    ("equipamento", FiltroRelatorio(escopo="brasil", equipamento=EQUIPAMENTO), (1, 3)),
    ("vazio", FiltroRelatorio(escopo="cnes", cnes="9999999"), ()),
    (
        "combinado",
        FiltroRelatorio(
            escopo="uf",
            uf="BA",
            ano_inicio=2025,
            ano_fim=2026,
            situacao="Em execução",
            programa="Programa matriz A",
            tipo_contratacao="Convênio",
            equipamento=EQUIPAMENTO,
        ),
        (1,),
    ),
]


@pytest.fixture
def carteira_matriz():
    with SessionLocal() as db:
        criar_carteira_matriz(db)
        try:
            yield db
        finally:
            db.rollback()


@pytest.mark.parametrize("formato", ["xlsx", "docx"])
@pytest.mark.parametrize("nivel", ["simplificado", "completo"])
@pytest.mark.parametrize("nome,filtro,esperados", CASOS, ids=[caso[0] for caso in CASOS])
def test_filtro_exporta_exatamente_os_convenios_esperados(carteira_matriz, formato, nivel, nome, filtro, esperados):
    # Given: carteira sintética, sem commit e com um registro de obra oculto.
    filtro = replace(filtro, busca=PREFIXO)
    # When: service/repositories/builders reais geram o arquivo final.
    conteudo = montar_relatorio(
        db=carteira_matriz, formato=formato, nivel=nivel, tipo_relatorio="instrumentos_repasse", filtro=filtro
    )
    # Then: o oráculo é a lista explícita, nunca a própria query testada.
    numeros = {f"900700{n}" for n in esperados}
    if formato == "xlsx":
        pasta = load_workbook(BytesIO(conteudo))
        obtidos = {str(linha[0].value) for linha in pasta["Convênios"].iter_rows(min_row=2)}
        assert obtidos == numeros
        assert pasta["Convênios"].freeze_panes == "A2"
        resumo = pasta["Resumo"]
        assert sum(
            linha[1].value
            for linha in resumo.iter_rows(min_row=2)
            if not str(linha[0].value).startswith("Instrumentos monitorados")
        ) == len(numeros)
    else:
        documento = Document(BytesIO(conteudo))
        # O oráculo desta matriz é a seção de Convênios, como a aba acima.
        # Acompanhamento interno usa filtros próprios e é certificado na
        # matriz cruzada; seus números não são Convênios extras no recorte.
        obtidos = set()
        for paragrafo in documento.paragraphs:
            estilo = paragrafo.style
            assert estilo is not None
            if estilo.name == "Heading 2" and paragrafo.text == "Acompanhamento interno no recorte":
                break
            if paragrafo.text.startswith("Número: "):
                obtidos.add(paragrafo.text.removeprefix("Número: "))
        assert obtidos == numeros
        assert documento.paragraphs[0].text.endswith(filtro.titulo())


def test_xlsx_preserva_identificadores_datas_zero_nulo_e_precisao(carteira_matriz):
    # Given / When
    conteudo = montar_relatorio(
        db=carteira_matriz,
        formato="xlsx",
        nivel="completo",
        tipo_relatorio="instrumentos_repasse",
        filtro=FiltroRelatorio(escopo="brasil", busca=PREFIXO),
    )
    aba = load_workbook(BytesIO(conteudo))["Convênios"]
    colunas = {celula.value: celula.column for celula in aba[1]}
    linhas = {aba.cell(n, colunas["Número"]).value: n for n in range(2, aba.max_row + 1)}

    def celula(numero, coluna):
        return aba.cell(linhas[numero], colunas[coluna])

    # Then: IDs não são números; valor monetário/data não são strings pt-BR.
    assert celula("9007001", "CNES").value == "0000701"
    assert celula("9007001", "CNPJ").value == "00000000000001"
    assert celula("9007001", "Valor global").value == 1234.56
    assert celula("9007001", "Valor global").number_format == '"R$" #,##0.00'
    assert celula("9007003", "Valor global").value == 0
    assert celula("9007004", "Valor global").value is None
    assert celula("9007001", "Valor pago ao fornecedor").value == 0
    assert celula("9007001", "Publicação").value == datetime(2025, 1, 31)
    assert celula("9007001", "Publicação").number_format == "DD/MM/YYYY"
    assert aba.auto_filter.ref == f"A1:X{aba.max_row}"


def test_xlsx_trata_texto_de_fonte_como_texto_nunca_formula(carteira_matriz):
    # Given: evidência inerte de injeção; não contém URL nem executa nada.
    carteira_matriz.query(Convenio).filter_by(numero="9007001").one().convenente_nome = "=1+1"
    carteira_matriz.flush()
    # When
    conteudo = montar_relatorio(
        db=carteira_matriz,
        formato="xlsx",
        nivel="completo",
        tipo_relatorio="instrumentos_repasse",
        filtro=FiltroRelatorio(escopo="cnes", cnes="0000701"),
    )
    aba = load_workbook(BytesIO(conteudo))["Convênios"]
    # Then
    assert aba["B2"].value == "=1+1"
    assert aba["B2"].data_type == "s"


@pytest.mark.parametrize("formato", ["xlsx", "docx"])
@pytest.mark.parametrize("nivel", ["simplificado", "completo"])
def test_gera_amostra_sintetica_para_conferencia_visual(carteira_matriz, tmp_path, formato, nivel):
    # Given / When: um instrumento com objeto longo, identificadores e centavos.
    conteudo = montar_relatorio(
        db=carteira_matriz,
        formato=formato,
        nivel=nivel,
        tipo_relatorio="instrumentos_repasse",
        filtro=FiltroRelatorio(escopo="cnes", cnes="0000701", busca="9007001"),
    )
    caminho = tmp_path / f"instrumentos-{nivel}.{formato}"
    caminho.write_bytes(conteudo)
    # Then: a amostra salva contém o instrumento esperado; render não é simulado.
    if formato == "xlsx":
        aba = load_workbook(caminho)["Convênios"]
        assert aba.max_row == 2
        assert aba["A2"].value == "9007001"
    else:
        documento = Document(caminho)
        assert [p.text for p in documento.paragraphs if p.text.startswith("Número: ")] == ["Número: 9007001"]
