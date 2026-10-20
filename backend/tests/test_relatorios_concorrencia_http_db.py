"""ASGI HTTP → sessões SQL distintas → services/builders reais em sobreposição."""

from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
from threading import Barrier
from zipfile import ZipFile

import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app
from tests.fixtures_relatorios_concorrencia import (
    CLIENTES,
    SENHA_SINTETICA,
    PedidoRelatorio,
)
from tests.fixtures_relatorios_concorrencia import (
    relatorios_concorrentes as relatorios_concorrentes,
)


def _solicitar(pedido: PedidoRelatorio, email: str, inicio: Barrier):
    ip = f"198.51.100.{next(CLIENTES)}"
    try:
        with TestClient(app, base_url="https://testserver", client=(ip, 40000)) as cliente:
            if pedido.autenticar:
                login = cliente.post("/auth/login", json={"email": email, "password": SENHA_SINTETICA})
                assert login.status_code == 200
                assert "access_token" not in login.json()
                assert cliente.cookies.get("sigeo_access") and cliente.cookies.get("sigeo_refresh")
            inicio.wait(timeout=30)
            return cliente.get(
                "/relatorios", params={"formato": pedido.formato, "nivel": pedido.nivel, **pedido.filtros}
            )
    except Exception:
        inicio.abort()
        raise


def _numeros_do_arquivo(formato: str, conteudo: bytes) -> frozenset[str]:
    with ZipFile(BytesIO(conteudo)) as pacote:
        assert pacote.testzip() is None
    if formato == "xlsx":
        pasta = load_workbook(BytesIO(conteudo))
        try:
            return frozenset(str(linha[0].value) for linha in pasta["Convênios"].iter_rows(min_row=2))
        finally:
            pasta.close()
    documento = Document(BytesIO(conteudo))
    return frozenset(p.text.removeprefix("Número: ") for p in documento.paragraphs if p.text.startswith("Número: "))


def _conferir_respostas(pedidos, respostas, *, admite_ocupado=False):
    for pedido, resposta in zip(pedidos, respostas, strict=True):
        if admite_ocupado and pedido.formato == "docx" and pedido.status == 200 and resposta.status_code == 503:
            assert resposta.json() == {
                "error": "Há outro relatório DOCX sendo gerado. Tente novamente em instantes.",
                "detail": None,
            }
            assert resposta.headers["content-type"] == "application/json"
            assert "content-disposition" not in resposta.headers
            continue
        assert resposta.status_code == pedido.status
        if pedido.status != 200:
            assert resposta.headers["content-type"] == "application/json"
            assert "content-disposition" not in resposta.headers
            assert set(resposta.json()) == {"error", "detail"}
            continue
        tipo = "spreadsheetml.sheet" if pedido.formato == "xlsx" else "wordprocessingml.document"
        assert resposta.headers["content-type"] == f"application/vnd.openxmlformats-officedocument.{tipo}"
        escopo = pedido.filtros["escopo"]
        nome = f"relatorio-instrumentos_repasse-{escopo}-{pedido.nivel}.{pedido.formato}"
        assert resposta.headers["content-disposition"] == f'attachment; filename="{nome}"'
        assert _numeros_do_arquivo(pedido.formato, resposta.content) == pedido.esperados


def _pedidos(formato: str, nivel: str):
    return [
        PedidoRelatorio(formato, nivel, {"escopo": "uf", "uf": "BA"}, frozenset({"9007001", "9007003", "9007005"})),
        PedidoRelatorio(formato, nivel, {"escopo": "uf", "uf": "PE"}, frozenset({"9007002"})),
        PedidoRelatorio(
            formato, nivel, {"escopo": "municipio", "uf": "SP", "municipio": "São Paulo"}, frozenset({"9007004"})
        ),
        PedidoRelatorio(formato, nivel, {"escopo": "cnes", "cnes": "9999999"}, frozenset()),
    ]


@pytest.mark.parametrize("formato", ["xlsx", "docx"])
@pytest.mark.parametrize("nivel", ["simplificado", "completo"])
def test_quatro_perfis_exportam_recortes_isolados_com_sessoes_sql_independentes(
    relatorios_concorrentes, formato, nivel
):
    # Given
    usuarios, observacoes, nome_banco, carteira = relatorios_concorrentes
    antes = carteira()
    pedidos = _pedidos(formato, nivel)
    inicio = Barrier(len(pedidos))
    # When
    with ThreadPoolExecutor(max_workers=len(pedidos)) as executor:
        tarefas = [
            executor.submit(_solicitar, pedido, email, inicio) for pedido, email in zip(pedidos, usuarios, strict=True)
        ]
        respostas = [tarefa.result(timeout=90) for tarefa in tarefas]
    # Then
    _conferir_respostas(pedidos, respostas, admite_ocupado=formato == "docx")
    assert any(resposta.status_code == 200 for resposta in respostas)
    assert len(observacoes) == len(pedidos)
    assert len({pid for pid, _ in observacoes}) == len(pedidos)
    assert {banco for _, banco in observacoes} == {nome_banco}
    assert carteira() == antes


def test_401_e_422_sobrepostos_nao_entregam_anexo_nem_afetam_os_downloads_validos(relatorios_concorrentes):
    # Given
    usuarios, _, _, carteira = relatorios_concorrentes
    antes = carteira()
    pedidos = [
        _pedidos("xlsx", "simplificado")[0],
        _pedidos("docx", "completo")[1],
        PedidoRelatorio("xlsx", "completo", {"escopo": "regiao"}, frozenset(), status=422),
        PedidoRelatorio("docx", "simplificado", {"escopo": "brasil"}, frozenset(), status=401, autenticar=False),
    ]
    inicio = Barrier(len(pedidos))
    # When
    with ThreadPoolExecutor(max_workers=len(pedidos)) as executor:
        tarefas = [
            executor.submit(_solicitar, pedido, email, inicio) for pedido, email in zip(pedidos, usuarios, strict=True)
        ]
        respostas = [tarefa.result(timeout=90) for tarefa in tarefas]
    # Then
    _conferir_respostas(pedidos, respostas)
    assert carteira() == antes
