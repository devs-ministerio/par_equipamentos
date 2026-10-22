"""Arquivos finais nos tetos reais; não é SLA, teste de RAM ou aceite visual."""

from io import BytesIO
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook
from sqlalchemy import insert

from app.db.base import SessionLocal
from app.services.relatorios import FiltroRelatorio, TipoRelatorio, montar_relatorio
from tests.fixtures_relatorios_limites import MODELOS, TETOS, linha_limite
from tests.fixtures_relatorios_merito import criar_execucao
from tests.fixtures_relatorios_monitoramento_limite import criar_monitorados_limite

CASOS = [
    ("convenios", "simplificado", "Convênios", 5000),
    ("convenios", "completo", "Convênios", 5000),
    ("propostas", "simplificado", "Propostas", 5000),
    ("propostas", "completo", "Propostas", 5000),
    ("macros", "simplificado", "Cobertura QA_LIMITE", 1000),
    ("municipios", "completo", "Cobertura QA_LIMITE", 10000),
    ("monitoramento", "simplificado", "Monitoramento", 500),
    ("monitoramento", "completo", "Monitoramento (timeline)", 500),
]


@pytest.fixture
def volume_arquivo(request):
    origem = request.param
    with SessionLocal() as db:
        try:
            if origem == "monitoramento":
                criar_monitorados_limite(db, quantidade=500)
            else:
                execucao = criar_execucao(db, "QA_LIMITE", "__qa_limite__", 2099, publicada=True)
                linhas = [
                    linha_limite(origem, indice, execucao.id, municipio=f"Município QA {indice:05}")
                    for indice in range(TETOS[origem])
                ]
                if origem in {"convenios", "propostas"}:
                    campo = "valor_global" if origem == "convenios" else "vl_global_proposta"
                    # Zero e ausência são contrastes explícitos, não preenchimento
                    # fictício de uma fonte real. Primeiro e último se distinguem.
                    for indice, linha in enumerate(linhas):
                        linha[campo] = 0 if indice % 2 == 0 else None
                db.execute(insert(MODELOS[origem]), linhas)
                db.flush()
            yield db
        finally:
            db.rollback()


def _esperados(origem: str, total: int) -> set[str]:
    if origem == "convenios":
        return {f"QA_LIMITE_{i:05}" for i in range(total)}
    if origem == "propostas":
        return {str(9_800_000 + i) for i in range(total)}
    if origem == "macros":
        return {f"Macro {i:05}" for i in range(total)}
    if origem == "municipios":
        return {f"Município QA {i:05}" for i in range(total)}
    return {f"QA_MON_{i:04}" for i in range(total)}


@pytest.mark.parametrize("volume_arquivo,nivel,aba,total", CASOS, indirect=["volume_arquivo"])
def test_xlsx_no_teto_contem_todo_o_recorte_sem_repeticao(volume_arquivo, nivel, aba, total, request):
    # Given
    origem = request.node.callspec.params["volume_arquivo"]
    tipo: TipoRelatorio = "analise_merito" if origem in {"macros", "municipios"} else "instrumentos_repasse"
    esperados = _esperados(origem, total)
    # When
    conteudo = montar_relatorio(
        db=volume_arquivo,
        formato="xlsx",
        nivel=nivel,
        tipo_relatorio=tipo,
        filtro=FiltroRelatorio(escopo="uf", uf="AC"),
    )
    # Then
    with ZipFile(BytesIO(conteudo)) as pacote:
        assert pacote.testzip() is None
    pasta = load_workbook(BytesIO(conteudo))
    try:
        folha = pasta[aba]
        linhas = list(folha.iter_rows(min_row=2, values_only=True))
        obtidos = [str(linha[0]) for linha in linhas]
        assert len(linhas) == total
        assert len(set(obtidos)) == total
        assert set(obtidos) == esperados
        assert min(esperados) in obtidos and max(esperados) in obtidos
        assert folha.freeze_panes == "A2"
        assert folha.auto_filter.ref.endswith(str(total + 1))
        if origem in {"convenios", "propostas"}:
            coluna = next(c.column for c in folha[1] if c.value == "Valor global")
            valores = [linha[coluna - 1] for linha in linhas]
            assert valores.count(0) == total // 2
            assert valores.count(None) == total // 2
    finally:
        pasta.close()
