"""HTTP real por proxy local, não ASGI/TestClient nem aceite de provedor."""

import json
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
from threading import Barrier
from time import perf_counter
from uuid import uuid4
from zipfile import ZipFile

import pytest
import requests
from docx import Document
from openpyxl import load_workbook

from tests.fixtures_relatorios_tcp import SENHA_TCP
from tests.fixtures_relatorios_tcp import servidor_relatorios_tcp as servidor_relatorios_tcp
from tests.numeros_relatorio import numeros_dos_blocos


def numeros(formato, conteudo):
    with ZipFile(BytesIO(conteudo)) as pacote:
        assert pacote.testzip() is None
    if formato == "docx":
        documento = Document(BytesIO(conteudo))
        return numeros_dos_blocos(documento.paragraphs)
    pasta = load_workbook(BytesIO(conteudo))
    try:
        return [r[0] for r in pasta["Convênios"].iter_rows(min_row=2, values_only=True)]
    finally:
        pasta.close()


@pytest.mark.parametrize("modo", ["misto", "docx", "xlsx", "docx-sequencial"])
def test_proxy_tcp_exporta_recortes_com_sessao_real_sem_mutar_carteira(servidor_relatorios_tcp, tmp_path, modo):
    # Given: quatro perfis, domínio real e isolamento por processo/SQL.
    base, emails, dominio, _recursos = servidor_relatorios_tcp
    antes = dominio()
    clientes = []
    try:
        for email in emails:
            cliente = requests.Session()
            clientes.append(cliente)
            resposta = cliente.post(base + "/auth/login", json={"email": email, "password": SENHA_TCP}, timeout=15)
            assert resposta.status_code == 200
            tem_sessao = bool(cliente.cookies.get("sigeo_access") and cliente.cookies.get("sigeo_refresh"))
            assert tem_sessao, "Login não emitiu os cookies esperados."
            assert any(c.name == "sigeo_refresh" and c.path == "/api/auth" for c in cliente.cookies)
            cliente.headers["X-CSRF-Token"] = resposta.json()["csrf_token"]
        cookie_anterior = clientes[0].cookies.get("sigeo_refresh")
        sem_csrf = requests.Session()
        try:
            sem_csrf.cookies.update(clientes[0].cookies)
            assert sem_csrf.post(base + "/auth/refresh", timeout=15).status_code == 403
        finally:
            sem_csrf.close()
        renovacao = clientes[0].post(base + "/auth/refresh", timeout=15)
        assert renovacao.status_code == 200
        rotacionou = bool(clientes[0].cookies.get("sigeo_refresh") != cookie_anterior)
        assert rotacionou, "Refresh não rotacionou."
        clientes[0].headers["X-CSRF-Token"] = renovacao.json()["csrf_token"]
        todos = {f"QA_CARGA_{n:04}" for n in range(500)}
        pedidos = [
            ("docx", {"escopo": "uf", "uf": "AC"}, todos),
            ("xlsx", {"escopo": "uf", "uf": "AC"}, todos),
            ("docx", {"escopo": "cnes", "cnes": "0000810"}, {"QA_CARGA_0000"}),
            ("xlsx", {"escopo": "uf", "uf": "PE"}, set()),
        ]
        trabalhadores = 1 if modo == "docx-sequencial" else 4
        if modo != "misto":
            formato_lote = "docx" if modo == "docx-sequencial" else modo
            pedidos = [(formato_lote, {"escopo": "uf", "uf": "AC"}, todos) for _ in range(4)]
        inicio = Barrier(4)

        def baixar(indice):
            formato, filtro, esperados = pedidos[indice]
            if trabalhadores > 1:
                inicio.wait(timeout=20)
            comeco = perf_counter()
            resposta = clientes[indice].get(
                base + "/relatorios",
                params={
                    "formato": formato,
                    "nivel": "completo",
                    "tipo_relatorio": "instrumentos_repasse",
                    **filtro,
                },
                headers={"X-Trace-Id": uuid4().hex},
                timeout=(3, 90),
            )
            fim = perf_counter()
            if resposta.status_code == 503:
                assert trabalhadores == 4 and formato == "docx"
                assert resposta.json() == {
                    "error": "Há outro relatório DOCX sendo gerado. Tente novamente em instantes.",
                    "detail": None,
                }
                assert "content-disposition" not in resposta.headers
                return {"formato": formato, "status": 503, "inicio": comeco, "fim": fim}
            assert resposta.status_code == 200
            assert f'.{formato}"' in resposta.headers["content-disposition"]
            mime = "wordprocessingml.document" if formato == "docx" else "spreadsheetml.sheet"
            assert resposta.headers["content-type"] == f"application/vnd.openxmlformats-officedocument.{mime}"
            obtidos = numeros(formato, resposta.content)
            assert len(obtidos) == len(esperados) and set(obtidos) == esperados
            return {
                "formato": formato,
                "status": 200,
                "registros": len(obtidos),
                "inicio": comeco,
                "fim": fim,
                "duracao_segundos": fim - comeco,
                "bytes": len(resposta.content),
            }

        # When: mesmos quatro clientes; controle sequencial não usa barreira.
        with ThreadPoolExecutor(max_workers=trabalhadores) as executor:
            resultados = list(executor.map(baixar, range(4)))
        assert any(r["status"] == 200 for r in resultados)
        if modo == "docx":
            assert any(r["status"] == 503 for r in resultados)
        if trabalhadores == 1:
            assert all(a["fim"] <= b["inicio"] for a, b in zip(resultados, resultados[1:]))
        # Then: erros não devolvem arquivo, domínio permanece byte-equivalente.
        for resposta, status in (
            (requests.get(base + "/relatorios", timeout=15), 401),
            (clientes[0].get(base + "/relatorios", params={"formato": "docx", "escopo": "regiao"}, timeout=15), 422),
        ):
            assert resposta.status_code == status
            assert "content-disposition" not in resposta.headers
            assert resposta.headers["content-type"] == "application/json"
        assert dominio() == antes
        (tmp_path / "tcp-resultados.json").write_text(
            json.dumps(
                {
                    "modo": modo,
                    "clientes_simultaneos": trabalhadores,
                    "pedidos": resultados,
                    "sobreposicao_clientes_observada": max(r["inicio"] for r in resultados)
                    < min(r["fim"] for r in resultados),
                }
            ),
            encoding="utf-8",
        )
    finally:
        for cliente in clientes:
            cliente.close()
