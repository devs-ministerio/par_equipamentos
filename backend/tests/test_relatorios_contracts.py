"""Contratos HTTP de GET /relatorios (Plan Mode docs/arquitetura/
planmode-relatorios-2026-09-25.md, Blocos 2, 4 e 5)."""

from __future__ import annotations

from fastapi.testclient import TestClient

from app.main import app
from app.services.relatorios_admissao import reservar_geracao_docx

client = TestClient(app)


def test_docx_ocupado_nao_entrega_anexo_nem_bloqueia_xlsx(headers_autenticados):
    # Given: reserva real ocupada, sem mockar Service/Repository/builder.
    with reservar_geracao_docx():
        # When: ambos os relatórios Word disputam a mesma capacidade.
        for tipo in ("instrumentos_repasse", "analise_merito"):
            resposta = client.get(
                "/relatorios", params={"formato": "docx", "tipo_relatorio": tipo}, headers=headers_autenticados
            )
            # Then: erro controlado, sem arquivo nem detalhe de infraestrutura.
            assert resposta.status_code == 503
            assert resposta.json() == {
                "error": "Há outro relatório DOCX sendo gerado. Tente novamente em instantes.",
                "detail": None,
            }
            assert "content-disposition" not in resposta.headers
        excel = client.get("/relatorios", params={"formato": "xlsx"}, headers=headers_autenticados)
        assert excel.status_code == 200
        assert excel.content.startswith(b"PK")
        invalido = client.get(
            "/relatorios",
            params={"formato": "docx", "tipo_relatorio": "analise_merito", "escopo": "cnes", "cnes": "0000810"},
            headers=headers_autenticados,
        )
        assert invalido.status_code == 422
        assert client.get("/relatorios", params={"formato": "docx"}).status_code == 401
    liberado = client.get("/relatorios", params={"formato": "docx"}, headers=headers_autenticados)
    assert liberado.status_code == 200
    assert liberado.content.startswith(b"PK")


def test_gerar_relatorio_exige_sessao():
    response = client.get("/relatorios", params={"formato": "xlsx"})
    assert response.status_code == 401


def test_gerar_relatorio_xlsx_brasil_default_e_instrumentos_repasse(headers_autenticados):
    response = client.get("/relatorios", params={"formato": "xlsx"}, headers=headers_autenticados)

    assert response.status_code == 200
    assert response.headers["content-type"] == ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert (
        'filename="relatorio-instrumentos_repasse-brasil-simplificado.xlsx"' in response.headers["content-disposition"]
    )
    assert len(response.content) > 0


def test_gerar_relatorio_docx_uf_completo(headers_autenticados):
    response = client.get(
        "/relatorios",
        params={"formato": "docx", "nivel": "completo", "escopo": "uf", "uf": "DF"},
        headers=headers_autenticados,
    )

    assert response.status_code == 200
    assert response.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    assert 'filename="relatorio-instrumentos_repasse-uf-completo.docx"' in response.headers["content-disposition"]


def test_gerar_relatorio_analise_merito_xlsx(headers_autenticados):
    response = client.get(
        "/relatorios",
        params={"formato": "xlsx", "tipo_relatorio": "analise_merito", "escopo": "uf", "uf": "DF"},
        headers=headers_autenticados,
    )
    assert response.status_code == 200
    assert 'filename="relatorio-analise_merito-uf-simplificado.xlsx"' in response.headers["content-disposition"]


def test_gerar_relatorio_analise_merito_escopo_cnes_e_422(headers_autenticados):
    response = client.get(
        "/relatorios",
        params={"formato": "xlsx", "tipo_relatorio": "analise_merito", "escopo": "cnes", "cnes": "9000001"},
        headers=headers_autenticados,
    )
    assert response.status_code == 422


def test_gerar_relatorio_periodo_chega_aos_registros_exportados(headers_autenticados, monkeypatch):
    from io import BytesIO

    from openpyxl import load_workbook

    from app.db.base import SessionLocal, get_db
    from tests.fixtures_relatorios_matriz import PREFIXO, criar_carteira_matriz

    # Given: sessão SQL real compartilhada com a requisição, revertida ao final.
    with SessionLocal() as db:
        criar_carteira_matriz(db)
        monkeypatch.setitem(app.dependency_overrides, get_db, lambda: db)
        params = {"formato": "xlsx", "escopo": "brasil", "busca": PREFIXO}
        controle = client.get("/relatorios", params=params, headers=headers_autenticados)
        assert controle.status_code == 200
        assert load_workbook(BytesIO(controle.content))["Convênios"].max_row == 6
        # When: o período precisa chegar à query, não apenas retornar HTTP 200.
        response = client.get(
            "/relatorios", params={**params, "ano_inicio": 2025, "ano_fim": 2025}, headers=headers_autenticados
        )
        # Then: arquivo final contém só os dois registros de 2025.
        assert response.status_code == 200
        aba = load_workbook(BytesIO(response.content))["Convênios"]
        assert [linha[0].value for linha in aba.iter_rows(min_row=2)] == ["9007001", "9007004"]


def test_gerar_relatorio_escopo_regiao_sem_valor_e_422(headers_autenticados):
    response = client.get("/relatorios", params={"formato": "xlsx", "escopo": "regiao"}, headers=headers_autenticados)
    assert response.status_code == 422


def test_gerar_relatorio_formato_invalido_e_422(headers_autenticados):
    response = client.get("/relatorios", params={"formato": "pdf"}, headers=headers_autenticados)
    assert response.status_code == 422


def test_gerar_relatorio_tipo_relatorio_invalido_e_422(headers_autenticados):
    response = client.get(
        "/relatorios", params={"formato": "xlsx", "tipo_relatorio": "outro"}, headers=headers_autenticados
    )
    assert response.status_code == 422
