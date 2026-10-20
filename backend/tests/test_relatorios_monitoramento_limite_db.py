"""Teto depois do recorte final, nunca antes de calcular a fase/período."""

from datetime import date
from io import BytesIO

import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.db.base import SessionLocal, get_db
from app.main import app
from app.repositories.limites_relatorio import LimiteConsultaExcedido
from app.services import monitoramento_instrumentos
from app.services.monitoramento_instrumentos import listar_instrumentos_monitorados
from app.services.relatorios import FiltroRelatorio, _instrumentos, montar_relatorio
from tests.fixtures_relatorios_monitoramento_limite import (
    concluir_monitorados,
    criar_monitorados_limite,
    ocultar_monitorados,
)


@pytest.fixture
def carteira_monitorada(monkeypatch):
    class Relogio(date):
        @classmethod
        def today(cls):
            return cls(2026, 10, 19)

    monkeypatch.setattr(monitoramento_instrumentos, "date", Relogio)
    with SessionLocal() as db:
        instrumentos = criar_monitorados_limite(db)
        try:
            yield db, instrumentos
        finally:
            db.rollback()


def test_excesso_de_monitorados_nao_e_lista_parcial(carteira_monitorada):
    # Given / When / Then
    db, _ = carteira_monitorada
    with pytest.raises(LimiteConsultaExcedido):
        _instrumentos(db, FiltroRelatorio(escopo="uf", uf="AC"))


def test_ui_preserva_o_limite_atual_sem_mudar_contrato(carteira_monitorada):
    # Given
    db, instrumentos = carteira_monitorada
    # When
    obtidos = listar_instrumentos_monitorados(db=db, ufs=["AC"])
    # Then
    assert [i.instrumento.nr_convenio for i in obtidos] == [i.nr_convenio for i in instrumentos[:500]]


def test_teto_exato_depois_de_excluir_obra_e_aceito(carteira_monitorada):
    # Given
    db, instrumentos = carteira_monitorada
    ocultar_monitorados(db, instrumentos[:1])
    # When
    obtidos = _instrumentos(db, FiltroRelatorio(escopo="uf", uf="AC"))
    # Then
    assert [i.instrumento.nr_convenio for i in obtidos] == [i.nr_convenio for i in instrumentos[1:]]


@pytest.mark.parametrize("recorte", ["municipio", "contratacao", "ano_historico", "ano_corrente", "obras"])
@pytest.mark.parametrize("formato", ["xlsx", "docx"])
@pytest.mark.parametrize("nivel", ["simplificado", "completo"])
def test_filtros_nao_perdem_monitorado_depois_do_primeiro_lote(carteira_monitorada, recorte, formato, nivel):
    # Given: somente o último registro pertence ao recorte final.
    db, instrumentos = carteira_monitorada
    ultimo = instrumentos[-1]
    filtro = FiltroRelatorio(escopo="uf", uf="AC")
    if recorte == "municipio":
        ultimo.municipio = "SAO PAULO"
        filtro = FiltroRelatorio(escopo="municipio", uf="AC", municipio="São Paulo")
    elif recorte == "contratacao":
        ultimo.tipo_contratacao = "TED"
        filtro = FiltroRelatorio(escopo="uf", uf="AC", tipo_contratacao="TED")
    elif recorte == "ano_historico":
        ultimo.ano_instrumento = 2020
        filtro = FiltroRelatorio(escopo="uf", uf="AC", ano_inicio=2020)
    elif recorte == "ano_corrente":
        concluir_monitorados(db, instrumentos[:-1])
        filtro = FiltroRelatorio(escopo="uf", uf="AC", ano_inicio=2026)
    else:
        ocultar_monitorados(db, instrumentos[:-1])
    db.flush()
    # When
    conteudo = montar_relatorio(
        db=db, formato=formato, nivel=nivel, tipo_relatorio="instrumentos_repasse", filtro=filtro
    )
    # Then: arquivo final completo, não apenas contador do resumo.
    if formato == "xlsx":
        pasta = load_workbook(BytesIO(conteudo))
        aba = pasta["Monitoramento" if nivel == "simplificado" else "Monitoramento (timeline)"]
        assert [row[0].value for row in aba.iter_rows(min_row=2)] == [ultimo.nr_convenio]
    else:
        documento = Document(BytesIO(conteudo))
        assert [p.text for p in documento.paragraphs if p.text.startswith("Número: ")] == [
            f"Número: {ultimo.nr_convenio}"
        ]


@pytest.mark.parametrize("formato", ["xlsx", "docx"])
def test_http_excesso_monitoramento_recusa_anexo(carteira_monitorada, formato, headers_autenticados, monkeypatch):
    # Given
    db, _ = carteira_monitorada
    monkeypatch.setitem(app.dependency_overrides, get_db, lambda: db)
    # When
    with TestClient(app) as client:
        resposta = client.get(
            "/relatorios", headers=headers_autenticados, params={"formato": formato, "escopo": "uf", "uf": "AC"}
        )
    # Then
    assert resposta.status_code == 422
    assert "Refine os filtros" in resposta.json()["error"]
    assert "content-disposition" not in resposta.headers
