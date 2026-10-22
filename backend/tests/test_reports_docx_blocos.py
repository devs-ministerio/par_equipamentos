"""Equivalência OOXML com python-docx, sem benchmark frágil por tempo."""

from io import BytesIO

import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches

from app.reports.docx_builder import (
    _nova_tabela,
    _novo_paragrafo,
    adicionar_paragrafo,
    adicionar_titulo,
    gerar_bytes,
)


def _documento_com_secao(cenario):
    documento = Document()
    documento.add_paragraph("Bloco anterior")
    if cenario != "padrao":
        secao = documento.add_section()
        secao.page_width = Inches(11)
        secao.left_margin = Inches(0.7)
        secao.right_margin = Inches(1.3)
        if cenario == "sem_dimensoes":
            secao.page_width = None
            secao.left_margin = None
            secao.right_margin = None
        if cenario == "sem_secao_final":
            corpo = documento._element.body
            corpo.remove(corpo[-1])
    return documento


@pytest.mark.parametrize("cenario", ["padrao", "segunda_secao", "sem_dimensoes", "sem_secao_final"])
def test_blocos_preservam_xml_da_biblioteca_e_geometria_da_secao_vigente(cenario):
    # Given: seções distintas; o fallback também mantém a semântica original.
    esperado = _documento_com_secao(cenario)
    obtido = _documento_com_secao(cenario)
    # When: intercalar blocos deve manter a ordem e o mesmo pai de documento.
    for indice in range(3):
        esperado.add_paragraph(f"Campo {indice}\t& <descrição>\nOutra linha")
        adicionar_paragrafo(obtido, f"Campo {indice}\t& <descrição>\nOutra linha")
        tabela_esperada = esperado.add_table(rows=1, cols=2)
        tabela_obtida = _nova_tabela(obtido, 2)
        for tabela in (tabela_esperada, tabela_obtida):
            tabela.cell(0, 0).text = f"Número {indice}"
            tabela.add_row().cells[1].text = "R$ 0"
        esperado.add_heading(f"Próximo {indice}", level=2)
        adicionar_titulo(obtido, f"Próximo {indice}", nivel=2)
    # Then: igualdade integral, não apenas quantidade/texto dos blocos.
    assert obtido._element.xml == esperado._element.xml
    reaberto = Document(BytesIO(gerar_bytes(obtido)))
    assert reaberto._element.xml == esperado._element.xml
    assert len(reaberto.tables) == 3
    assert reaberto.tables[2].cell(1, 1).text == "R$ 0"


@pytest.mark.parametrize("nivel", range(10))
def test_titulo_preserva_estilo_e_texto_em_todos_os_niveis(nivel):
    # Given
    esperado, obtido = Document(), Document()
    # When
    esperado.add_heading("Título & acentuação\tTeste", level=nivel)
    adicionar_titulo(obtido, "Título & acentuação\tTeste", nivel=nivel)
    # Then
    assert obtido._element.xml == esperado._element.xml


@pytest.mark.parametrize("nivel", [-1, 10])
def test_titulo_invalido_nao_insere_bloco(nivel):
    # Given
    documento = Document()
    anterior = documento._element.xml
    # When / Then
    with pytest.raises(ValueError, match="level must be in range 0-9"):
        adicionar_titulo(documento, "Não deve existir", nivel=nivel)
    assert documento._element.xml == anterior


def test_paragrafo_vazio_sem_secao_final_preserva_api_da_biblioteca():
    # Given
    esperado, obtido = Document(), Document()
    for documento in (esperado, obtido):
        corpo = documento._element.body
        corpo.remove(corpo[-1])
    # When
    esperado.add_paragraph()
    paragrafo = _novo_paragrafo(obtido)
    paragrafo.add_run("Texto posterior").bold = True
    esperado.paragraphs[0].add_run("Texto posterior").bold = True
    # Then
    assert obtido._element.xml == esperado._element.xml
    assert obtido._element.body[-1].tag == qn("w:p")
