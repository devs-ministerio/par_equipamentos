"""app/reports/xlsx_builder.py -- builder generico de planilha (Plan Mode
docs/arquitetura/planmode-relatorios-2026-09-25.md, Bloco 1)."""

from __future__ import annotations

from datetime import date
from io import BytesIO

from openpyxl import load_workbook

from app.reports.xlsx_builder import colorir_variacao, escrever_aba_tabela, gerar_bytes, nova_pasta


def test_nova_pasta_nao_tem_aba_fantasma():
    pasta = nova_pasta()
    assert pasta.sheetnames == []


def test_escrever_aba_tabela_grava_cabecalho_e_linhas():
    pasta = nova_pasta()
    escrever_aba_tabela(pasta, "Cobertura", ["UF", "Cobertura %"], [["SP", 82.5], ["BA", 41.0]])

    reaberta = load_workbook(BytesIO(gerar_bytes(pasta)))
    aba = reaberta["Cobertura"]
    assert [c.value for c in aba[1]] == ["UF", "Cobertura %"]
    assert [c.value for c in aba[2]] == ["SP", 82.5]
    assert [c.value for c in aba[3]] == ["BA", 41.0]
    assert aba.freeze_panes == "A2"
    assert aba.auto_filter.ref == "A1:B3"


def test_titulo_de_aba_maior_que_31_caracteres_e_truncado():
    pasta = nova_pasta()
    aba = escrever_aba_tabela(pasta, "Um titulo de aba absurdamente comprido", [], [])
    assert len(aba.title) <= 31


def test_aba_sem_linhas_nao_define_autofilter():
    pasta = nova_pasta()
    aba = escrever_aba_tabela(pasta, "Vazia", ["A"], [])
    assert aba.auto_filter.ref is None


def test_colorir_variacao_positiva_fica_verde():
    pasta = nova_pasta()
    aba = escrever_aba_tabela(pasta, "Serie", ["Ano", "Var%"], [[2025, 10.0]])
    colorir_variacao(aba, coluna=2, linha=2, valor=10.0)
    assert aba.cell(row=2, column=2).font.color.rgb == "0015803D"


def test_colorir_variacao_negativa_fica_vermelha():
    pasta = nova_pasta()
    aba = escrever_aba_tabela(pasta, "Serie", ["Ano", "Var%"], [[2025, -5.0]])
    colorir_variacao(aba, coluna=2, linha=2, valor=-5.0)
    assert aba.cell(row=2, column=2).font.color.rgb == "00B91C1C"


def test_colorir_variacao_none_nao_pinta():
    pasta = nova_pasta()
    aba = escrever_aba_tabela(pasta, "Serie", ["Ano", "Var%"], [[2020, None]])
    colorir_variacao(aba, coluna=2, linha=2, valor=None)
    assert aba.cell(row=2, column=2).font.color.type == "theme"


def test_colunas_moeda_aplicam_number_format_mantendo_valor_cru():
    pasta = nova_pasta()
    aba = escrever_aba_tabela(
        pasta, "Convênios", ["Número", "Valor global"], [["123", 1234.5]], colunas_moeda=["Valor global"]
    )
    celula = aba.cell(row=2, column=2)
    assert celula.value == 1234.5
    assert celula.number_format == '"R$" #,##0.00'
    assert aba.cell(row=2, column=1).number_format == "General"


def test_colunas_data_aplicam_number_format_mantendo_valor_cru():
    pasta = nova_pasta()
    aba = escrever_aba_tabela(
        pasta, "Convênios", ["Número", "Publicação"], [["123", date(2026, 9, 26)]], colunas_data=["Publicação"]
    )
    celula = aba.cell(row=2, column=2)
    assert celula.value == date(2026, 9, 26)
    assert celula.number_format == "DD/MM/YYYY"


def test_coluna_moeda_inexistente_no_cabecalho_nao_quebra():
    pasta = nova_pasta()
    escrever_aba_tabela(pasta, "Vazia", ["A"], [["x"]], colunas_moeda=["Não existe"])


def test_texto_externo_nao_vira_formula_e_preserva_valor_original():
    # Arrange: diferentes prefixos perigosos, inclusive após whitespace.
    entradas = ["=1+1", "+SUM(1,2)", "-1+2", "@SUM(1,2)", "\t=2+2", "normal"]
    pasta = nova_pasta()

    # Act
    escrever_aba_tabela(pasta, "Propostas", ["Proponente", "Valor"], [[texto, 10] for texto in entradas])
    aba = load_workbook(BytesIO(gerar_bytes(pasta)))["Propostas"]

    # Assert
    for indice, texto in enumerate(entradas, start=2):
        assert aba.cell(indice, 1).value == texto
        assert aba.cell(indice, 1).data_type == "s"
        assert aba.cell(indice, 2).data_type == "n"
    assert aba.print_area == "'Propostas'!$A$1:$B$7"
