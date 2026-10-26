"""app/reports/docx_builder.py -- builder generico de documento (Plan Mode
docs/arquitetura/planmode-relatorios-2026-09-25.md, Bloco 1)."""

from __future__ import annotations

from datetime import date, datetime
from io import BytesIO

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

from app.reports.docx_builder import (
    adicionar_cabecalho_institucional,
    adicionar_fonte,
    adicionar_legenda_tabela,
    adicionar_paragrafo,
    adicionar_paragrafo_rotulo_valor,
    adicionar_tabela,
    adicionar_titulo,
    escrever_variacao,
    gerar_bytes,
    linhas_como_texto,
    novo_documento,
)


@pytest.mark.parametrize("total", [0, 1, 3, 5, 6, 8])
def test_tabelas_paginam_com_cabecalho_repetido_sem_partir_linhas(total):
    # Arrange / Act: tabela curta fica junta; tabela longa continua paginável.
    documento = novo_documento("Relatório")
    adicionar_tabela(documento, ["Marco", "Observação"], [[f"Marco {n}", "Descrição"] for n in range(total)])
    tabela = Document(BytesIO(gerar_bytes(documento))).tables[0]
    # Assert: propriedades sobrevivem ao arquivo final, não só ao objeto em memória.
    propriedades_cabecalho = tabela.rows[0]._tr.trPr
    assert propriedades_cabecalho is not None
    assert propriedades_cabecalho.find(qn("w:tblHeader")) is not None
    for linha in tabela.rows:
        propriedades_linha = linha._tr.trPr
        assert propriedades_linha is not None
        assert propriedades_linha.find(qn("w:cantSplit")) is not None
    # Tabela sem registros ganha uma linha de aviso, que também conta na paginação.
    registros = max(total, 1)
    for indice, linha in enumerate(tabela.rows):
        esperado = True if indice < registros and (registros + 1 <= 6 or indice == 0) else None
        assert linha.cells[0].paragraphs[0].paragraph_format.keep_with_next is esperado
    assert tabela.rows[-1].cells[0].paragraphs[0].paragraph_format.keep_with_next is None


def test_novo_documento_grava_titulo_como_heading():
    documento = novo_documento("Relatorio de Convenios")
    reaberto = Document(BytesIO(gerar_bytes(documento)))
    assert reaberto.paragraphs[0].text == "Relatorio de Convenios"
    estilo = reaberto.paragraphs[0].style
    assert estilo is not None and estilo.name == "Title"


def test_adicionar_titulo_usa_o_nivel_pedido():
    documento = novo_documento("Capa")
    adicionar_titulo(documento, "Convenios", nivel=1)
    reaberto = Document(BytesIO(gerar_bytes(documento)))
    assert reaberto.paragraphs[1].text == "Convenios"
    estilo = reaberto.paragraphs[1].style
    assert estilo is not None and estilo.name == "Heading 1"


def test_adicionar_paragrafo_grava_texto():
    documento = novo_documento("Capa")
    adicionar_paragrafo(documento, "Filtro: Regiao Nordeste")
    reaberto = Document(BytesIO(gerar_bytes(documento)))
    # O último parágrafo é o logotipo do SISGEO, centralizado ao fim do documento.
    assert reaberto.paragraphs[-2].text == "Filtro: Regiao Nordeste"


def test_logotipo_do_sisgeo_fecha_o_documento_centralizado():
    reaberto = Document(BytesIO(gerar_bytes(novo_documento("Capa"))))
    ultimo = reaberto.paragraphs[-1]
    assert ultimo.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert len(ultimo.runs[0]._r.xpath(".//pic:pic")) == 1


def test_adicionar_tabela_grava_cabecalho_em_negrito_e_linhas():
    documento = novo_documento("Capa")
    tabela = adicionar_tabela(documento, ["UF", "Convenios"], [["MG", 12], ["BA", 7]])
    assert len(tabela.rows) == 3
    assert tabela.rows[0].cells[0].text == "UF"
    assert tabela.rows[0].cells[0].paragraphs[0].runs[0].bold is True
    assert tabela.rows[1].cells[0].text == "MG"
    assert tabela.rows[2].cells[1].text == "7"


def test_adicionar_tabela_valor_none_vira_celula_vazia():
    documento = novo_documento("Capa")
    tabela = adicionar_tabela(documento, ["UF", "Nota"], [["MG", None]])
    assert tabela.rows[1].cells[1].text == ""


@pytest.mark.parametrize("pesos", [[1], [1, 0], [1, -1], [float("nan"), 1]])
def test_tabela_recusa_larguras_incoerentes_sem_alterar_documento(pesos):
    # Arrange
    documento = novo_documento("Relatório")
    # Act / Assert
    with pytest.raises(ValueError, match="peso positivo"):
        adicionar_tabela(documento, ["Marco", "UF"], [], pesos_colunas=pesos)
    assert len(documento.tables) == 0


def test_tabela_personalizada_preserva_textos_e_largura_util():
    # Arrange
    documento = novo_documento("Relatório")
    adicionar_tabela(documento, ["Descrição", "UF"], [["Texto extenso", "BA"]], pesos_colunas=[3, 1], fonte_pt=11)
    # Act
    reaberto = Document(BytesIO(gerar_bytes(documento)))
    tabela = reaberto.tables[0]
    # Assert: layout, conteúdo e paginação sobreviveram ao arquivo final.
    secao = reaberto.sections[-1]
    assert secao.page_width is not None
    assert secao.left_margin is not None
    assert secao.right_margin is not None
    assert (
        abs(sum(c.width for c in tabela.columns) - (secao.page_width - secao.left_margin - secao.right_margin)) < 1500
    )
    assert tabela.columns[0].width == pytest.approx(tabela.columns[1].width * 3, abs=1500)
    assert tabela.autofit is False
    assert [c.text for c in tabela.rows[1].cells] == ["Texto extenso", "BA"]
    assert tabela.rows[0].cells[0].paragraphs[0].runs[0].bold is True
    assert all(r.font.size.pt == 11 for row in tabela.rows for c in row.cells for p in c.paragraphs for r in p.runs)
    assert tabela.rows[0]._tr.xpath("./w:trPr/w:tblHeader")


def test_escrever_variacao_positiva_fica_verde():
    documento = novo_documento("Capa")
    paragrafo = documento.add_paragraph()
    run = escrever_variacao(paragrafo, 12.3)
    assert run.text == "+12.3%"
    assert str(run.font.color.rgb) == "15803D"


def test_escrever_variacao_negativa_fica_vermelha():
    documento = novo_documento("Capa")
    paragrafo = documento.add_paragraph()
    run = escrever_variacao(paragrafo, -4.0)
    assert run.text == "-4.0%"
    assert str(run.font.color.rgb) == "B91C1C"


def test_escrever_variacao_none_escreve_traco_sem_cor():
    documento = novo_documento("Capa")
    paragrafo = documento.add_paragraph()
    run = escrever_variacao(paragrafo, None)
    assert run.text == "--"
    assert run.font.color.rgb is None


def test_adicionar_cabecalho_institucional_grava_timbre_e_aviso_automatico():
    documento = novo_documento("Capa")
    adicionar_cabecalho_institucional(documento, gerado_em=datetime(2026, 9, 26, 10, 0))
    reaberto = Document(BytesIO(gerar_bytes(documento)))
    textos = [p.text for p in reaberto.paragraphs]
    assert "Ministério da Saúde" in textos
    assert any("DECAN" in t for t in textos)
    assert any("gerado automaticamente pelo SISGEO" in t for t in textos)


def test_adicionar_paragrafo_rotulo_valor_negrito_so_no_rotulo():
    documento = novo_documento("Capa")
    adicionar_paragrafo_rotulo_valor(documento, "Objeto", "Aquisição de equipamento")
    paragrafo = documento.paragraphs[-1]
    assert paragrafo.text == "Objeto: Aquisição de equipamento"
    assert paragrafo.runs[0].bold is True
    assert paragrafo.runs[1].bold is not True


def test_adicionar_legenda_tabela_numera_e_usa_o_texto():
    documento = novo_documento("Capa")
    adicionar_legenda_tabela(documento, 3, "Convênios firmados no recorte.")
    assert documento.paragraphs[-1].text == "Tabela 3. Convênios firmados no recorte."


def test_adicionar_fonte_prefixa_com_fonte():
    documento = novo_documento("Capa")
    adicionar_fonte(documento, "SISGEO.")
    assert documento.paragraphs[-1].text == "Fonte: SISGEO."


def test_linhas_como_texto_formata_moeda_e_data_por_nome_de_coluna():
    colunas = ["Convênio", "Valor global", "Publicação"]
    linhas: list[list[object]] = [["123", 1234.5, date(2026, 9, 26)], ["456", None, None]]
    resultado = linhas_como_texto(colunas, linhas, colunas_moeda=["Valor global"], colunas_data=["Publicação"])
    assert resultado == [
        ["123", "R$ 1.234,50", "26/09/2026"],
        ["456", "—", "—"],
    ]
