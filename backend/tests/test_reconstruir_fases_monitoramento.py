from datetime import datetime

import pytest
from openpyxl import Workbook

from scripts.complementar_persus_monitoramento import ler_controle
from scripts.reconstruir_fases_monitoramento import fase_para_evento


def test_fase_para_evento_prefere_vinculo_ja_persistido():
    assert fase_para_evento("cronograma_entrega", "fase_equipamento_entregue") == "fase_equipamento_entregue"


def test_fase_para_evento_inferre_somente_marcos_inequivocos():
    assert fase_para_evento("cronograma_comissionamento", None) == "fase_comissionamento"
    assert fase_para_evento("cronograma_trp", None) is None


def criar_controle_com_alias(arquivo, *, duplicar_cnes):
    livro = Workbook()
    obras = livro.active
    assert obras is not None
    obras.title = "Obras"
    obras.append(["Cabeçalho sintético"] * 13)
    obras.append(["Coluna sintética"] * 13)
    obra = [
        "Obra QA",
        "00000.000001/2025-00",
        "0000001",
        "DF",
        "Cidade QA",
        "Hospital QA - obras",
        None,
        None,
        None,
        None,
        0,
        None,
        0,
    ]
    obras.append(obra)
    equipamentos = livro.create_sheet("Equipamentos")
    equipamentos.append(["Coluna sintética"] * 21)
    equipamento: list[str | int | datetime | None] = [None] * 21
    equipamento[1:5] = ["0000001", "DF", "Cidade QA", "Hospital QA - equipamentos"]
    equipamento[6] = 2025
    equipamento[15] = datetime(2025, 10, 10)
    equipamentos.append(equipamento)
    if duplicar_cnes:
        obras.append(obra)
        equipamentos.append(equipamento)
    inauguracao = livro.create_sheet("Inauguração")
    inauguracao.append(["Coluna sintética"] * 11)
    inauguracao.append(["DF", "AL", "0000001", "Cidade QA", "Hospital QA", None, None, None, None, None, None])
    livro.save(arquivo)
    livro.close()


@pytest.mark.parametrize("duplicar_cnes", [False, True], ids=["cnes-unico", "cnes-ambiguo"])
def test_controle_so_usa_cnes_unico_quando_unidade_tem_alias_entre_abas(tmp_path, duplicar_cnes):
    # Given: três nomes diferentes, mesmo CNES sintético; sem planilha operacional.
    arquivo = tmp_path / "controle-sintetico.xlsx"
    criar_controle_com_alias(arquivo, duplicar_cnes=duplicar_cnes)
    # When: parser real, sem substituir regra de associação.
    registros = ler_controle(arquivo)
    # Then: duplicidade não escolhe arbitrariamente uma linha complementar.
    assert len(registros) == 1
    registro = registros[0]
    assert registro.cnes_informado == "0000001"
    if duplicar_cnes:
        assert registro.datas_equipamento == {}
        assert registro.nup is None
    else:
        assert registro.datas_equipamento["cronograma_comissionamento"].isoformat() == "2025-10-10"
        assert registro.nup == "00000.000001/2025-00"
