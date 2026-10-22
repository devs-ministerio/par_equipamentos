"""Integridade DOCX nos tetos sintéticos, não aceite visual ou SLA nacional."""

import json
from io import BytesIO
from time import perf_counter
from zipfile import ZipFile

import pytest
from docx import Document

from app.services.relatorios import FiltroRelatorio, TipoRelatorio, montar_relatorio
from tests.numeros_relatorio import numeros_dos_blocos
from tests.test_relatorios_volume_arquivos_db import _esperados
from tests.test_relatorios_volume_arquivos_db import volume_arquivo as volume_arquivo

CASOS = [
    ("convenios", "completo", 5000),
    ("propostas", "completo", 5000),
    ("macros", "simplificado", 1000),
    ("municipios", "completo", 10000),
    ("monitoramento", "simplificado", 500),
    ("monitoramento", "completo", 500),
]


@pytest.mark.parametrize("volume_arquivo,nivel,total", CASOS, indirect=["volume_arquivo"])
def test_docx_no_teto_preserva_todos_os_blocos_ou_linhas(volume_arquivo, nivel, total, request, tmp_path):
    # Given
    origem = request.node.callspec.params["volume_arquivo"]
    tipo: TipoRelatorio = "analise_merito" if origem in {"macros", "municipios"} else "instrumentos_repasse"
    esperados = _esperados(origem, total)
    # When
    inicio = perf_counter()
    conteudo = montar_relatorio(
        db=volume_arquivo,
        formato="docx",
        nivel=nivel,
        tipo_relatorio=tipo,
        filtro=FiltroRelatorio(escopo="uf", uf="AC"),
    )
    geracao_segundos = perf_counter() - inicio
    # Then
    with ZipFile(BytesIO(conteudo)) as pacote:
        assert pacote.testzip() is None
    documento = Document(BytesIO(conteudo))
    if origem in {"macros", "municipios"}:
        cabecalho = "Macrorregião" if origem == "macros" else "Município"
        tabelas = [
            tabela
            for tabela in documento.tables
            if tabela.rows[0].cells[0].text == cabecalho and len(tabela.rows) == total + 1
        ]
        assert len(tabelas) == 1
        linhas = iter(tabelas[0].rows)
        next(linhas)  # Cabeçalho não é registro.
        obtidos = [linha.cells[0].text for linha in linhas]
    elif origem == "propostas":
        obtidos = [
            p.text.rsplit(" -- Proposta ", 1)[1]
            for p in documento.paragraphs
            if " -- Proposta " in p.text and p.style is not None and p.style.name == "Heading 3"
        ]
    else:
        obtidos = numeros_dos_blocos(documento.paragraphs)
    assert len(obtidos) == total
    assert len(set(obtidos)) == total
    assert set(obtidos) == esperados
    for tabela in documento.tables:
        linhas = iter(tabela.rows)
        assert next(linhas)._tr.xpath("./w:trPr/w:tblHeader")
        assert all(linha._tr.xpath("./w:trPr/w:cantSplit") for linha in tabela.rows)
    (tmp_path / "medicao-docx.json").write_text(
        json.dumps(
            {
                "origem": origem,
                "nivel": nivel,
                "registros": total,
                "geracao_segundos": geracao_segundos,
                "validacao_segundos": perf_counter() - inicio - geracao_segundos,
                "bytes": len(conteudo),
            }
        ),
        encoding="utf-8",
    )
